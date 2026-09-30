"""Shared styles, layout helpers, lists registry and the formula linter."""
import re

from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- design tokens
FONT = "Arial"
ACCENT = "0F5E5E"
ACCENT_LIGHT = "E3F1F1"
ACCENT_MID = "9CCFCF"
INPUT = "FFF7CC"
AUTO_BG = "F3F4F6"
WHITE = "FFFFFF"
GREY_TEXT = "6B7280"
DARK = "1F2937"
LINE = "D1D5DB"
GOOD = ("D1FAE5", "065F46")
WARN = ("FEF3C7", "92400E")
BAD = ("FEE2E2", "991B1B")
TAB_TEAL, TAB_GREEN, TAB_ORANGE, TAB_PURPLE, TAB_GREY = "0F5E5E", "2E7D32", "E07B00", "6A1B9A", "9CA3AF"

DATE = "dd mmm yyyy"
MONEY = '#,##0;[Red]-#,##0;"–"'
PCT = "0%"
INT = "0"
NUM = "#,##0"

# Status strings shared by formulas, conditional formats and the test model.
ICON_OK, ICON_WARN, ICON_BAD = "🟢", "🟡", "🔴"
FLAG_OVERPAY = "⚠️ Paying ahead of work"
ST_DECIDE = "Decide now"
ST_WAITING = "Waiting for decision"
ST_ORDER_NOW = "Order now"
ST_ORDER_LATE = "⚠️ Order late — install at risk"
ST_ORDERED = "Ordered"
ST_DELIV_LATE = "⚠️ Delivery after install date"
ST_DELIVERED = "Delivered"
ST_SELECTED = "Selected"

# Sheet names (in order).
S_START = "Start Here"
S_DASH = "Dashboard"
S_WEEK = "This Week"
S_BUDGET = "Budget"
S_QUOTES = "Quotes"
S_CON = "Contractors"
S_PAY = "Payments"
S_CO = "Change Orders"
S_TL = "Timeline"
S_ROOMS = "Rooms"
S_SEL = "Selections & Orders"
S_ISS = "Issues & Punch List"
S_VAULT = "Vault"
S_WAR = "Warranty & Maintenance"
S_ENG = "Engine"
S_LISTS = "Lists"
SHEET_ORDER = [S_START, S_DASH, S_WEEK, S_BUDGET, S_QUOTES, S_CON, S_PAY, S_CO, S_TL,
               S_ROOMS, S_SEL, S_ISS, S_VAULT, S_WAR, S_ENG, S_LISTS]
# Lite edition: the core money + schedule loop only (other tabs are not in the file at all).
LITE_SHEETS = [S_START, S_DASH, S_BUDGET, S_CON, S_PAY, S_TL, S_ISS, S_ENG, S_LISTS]


def q(sheet):
    """Quoted sheet name for formulas."""
    return "'" + sheet.replace("'", "''") + "'"


def col_idx(letter):
    return column_index_from_string(letter)


def col_let(idx):
    return get_column_letter(idx)


def col_add(letter, n):
    return col_let(col_idx(letter) + n)


def pri_from_due(d):
    """Priority from a due date cell: 1 overdue, 2 within alert window, 3 within 2x window."""
    return (f'IF({d}="","",IF({d}<AsOf,1,IF({d}<=AsOf+AlertDays,2,'
            f'IF({d}<=AsOf+2*AlertDays,3,""))))')


def money_txt(expr):
    """Text like '€ 1,250' for alert texts (FIXED follows the user's locale)."""
    return f'CurSym&" "&FIXED({expr},0)'


# ---------------------------------------------------------------- styles
_side = Side(style="thin", color=LINE)
BORDER_ALL = Border(left=_side, right=_side, top=_side, bottom=_side)
BORDER_BOTTOM = Border(bottom=Side(style="medium", color=ACCENT))
NO_BORDER = Border()


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def font(size=10, bold=False, color=DARK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)


def style(cell, kind="auto", fmt=None, bold=False, size=10, align=None, wrap=False,
          color=None, border=True, italic=False):
    """Apply one of the standard cell looks."""
    if kind == "input":
        cell.fill = fill(INPUT)
        cell.protection = UNLOCKED
        cell.font = font(size, bold, color or DARK, italic)
    elif kind == "auto":
        cell.fill = fill(WHITE)
        cell.protection = LOCKED
        cell.font = font(size, bold, color or DARK, italic)
    elif kind == "autogrey":
        cell.fill = fill(AUTO_BG)
        cell.protection = LOCKED
        cell.font = font(size, bold, color or DARK, italic)
    elif kind == "header":
        cell.fill = fill(ACCENT)
        cell.protection = LOCKED
        cell.font = font(size, True, WHITE)
        wrap = True
        align = align or "center"
    elif kind == "band":
        cell.fill = fill(ACCENT_LIGHT)
        cell.protection = LOCKED
        cell.font = font(size, True, color or ACCENT)
    elif kind == "label":
        cell.protection = LOCKED
        cell.font = font(size, bold, color or DARK, italic)
        border = False
    elif kind == "note":
        cell.protection = LOCKED
        cell.font = font(9, bold, color or GREY_TEXT, True)
        border = False
    elif kind == "hidden":
        cell.protection = LOCKED
        cell.font = font(9, False, GREY_TEXT)
        border = False
    if border and kind not in ("label", "note", "hidden"):
        cell.border = BORDER_ALL
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)


def put(ws, ref, value, kind="auto", fmt=None, **kw):
    c = ws[ref]
    c.value = value
    style(c, kind, fmt, **kw)
    return c


def comment(cell, text, width=260, height=110):
    cm = Comment(text, "Renovation OS")
    cm.width = width
    cm.height = height
    cell.comment = cm


def title(ws, text, desc, width_to="K"):
    ws["B1"].value = text
    ws["B1"].font = Font(name=FONT, size=18, bold=True, color=ACCENT)
    ws["B1"].alignment = Alignment(vertical="center")
    ws["B2"].value = desc
    ws["B2"].font = Font(name=FONT, size=10, italic=True, color=GREY_TEXT)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 20
    ws.column_dimensions["A"].width = 2.5


def section(ws, ref, text, span_to=None, size=11):
    c = ws[ref]
    c.value = text
    c.font = font(size, True, ACCENT)
    c.border = BORDER_BOTTOM
    c.alignment = Alignment(vertical="center")
    if span_to:
        row = c.row
        for ci in range(c.column + 1, col_idx(span_to) + 1):
            ws.cell(row=row, column=ci).border = BORDER_BOTTOM


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def hide_cols(ws, letters):
    for L in letters:
        ws.column_dimensions[L].hidden = True


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.formatCells = True
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    p.sort = True
    p.autoFilter = True
    p.insertRows = True
    p.deleteRows = True


def add_cf(ws, rng, formula, colors, bold=False, stop=False):
    bg, fg = colors
    rule = FormulaRule(formula=[formula], fill=fill(bg), font=Font(name=FONT, color=fg, bold=bold),
                       stopIfTrue=stop)
    ws.conditional_formatting.add(rng, rule)


def add_cf_fill(ws, rng, formula, bg, fg=None, stop=False):
    kw = {"fill": fill(bg), "stopIfTrue": stop}
    if fg:
        kw["font"] = Font(name=FONT, color=fg)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], **kw))


def level_cf(ws, rng, level_ref):
    """Colour a status range from a numeric level cell (0 good, 1 attention, 2 urgent)."""
    add_cf(ws, rng, f"{level_ref}=2", BAD, bold=True)
    add_cf(ws, rng, f"{level_ref}=1", WARN, bold=True)
    add_cf(ws, rng, f"{level_ref}=0", GOOD)


def define(wb, name, sheet, ref):
    dn = DefinedName(name, attr_text=f"{q(sheet)}!{ref}")
    wb.defined_names[name] = dn


# ---------------------------------------------------------------- tables
class Table:
    """A fixed-size list: header row + n data rows; columns addressed by key."""

    def __init__(self, sheet, header_row, first, n, cols):
        self.sheet = sheet
        self.header_row = header_row
        self.first = first
        self.n = n
        self.last = first + n - 1
        self.cols = {}
        cur = "B"
        for c in cols:
            if isinstance(c, tuple):
                key, cur = c
            else:
                key = c
            self.cols[key] = cur
            cur = col_add(cur, 1)

    def L(self, key):
        return self.cols[key]

    def c(self, key, r):
        return f"{self.cols[key]}{r}"

    def local(self, key):
        L = self.cols[key]
        return f"${L}${self.first}:${L}${self.last}"

    def a(self, key):
        return f"{q(self.sheet)}!{self.local(key)}"

    def f(self, template, r):
        """Fill {key} placeholders with cell refs on row r; {r} is the row number."""
        m = {k: f"{v}{r}" for k, v in self.cols.items()}
        m["r"] = r
        m["r1"] = r - 1
        return template.format_map(m)

    def rows(self):
        return range(self.first, self.last + 1)


def write_table(ws, T, spec, data=None, dv=None, row_height=20, id_start=1):
    """Write headers, formulas, inputs and demo data for a Table.

    spec: list of dicts with key, header, width, kind in {'in','auto','id','hid','in_f'},
          fmt, list (dropdown list name), f (formula template), note (comment), wrap, align,
          default (value or formula written to input cells when no data).
    """
    data = data or []
    hr = T.header_row
    ws.row_dimensions[hr].height = 32
    for s in spec:
        L = T.L(s["key"])
        h = ws[f"{L}{hr}"]
        h.value = s.get("header", "")
        style(h, "header" if s["kind"] != "hid" else "hidden")
        if s.get("note"):
            comment(h, s["note"])
        if s.get("width"):
            ws.column_dimensions[L].width = s["width"]
        if s["kind"] == "hid":
            ws.column_dimensions[L].hidden = True
        if s.get("list") and dv is not None:
            dv.add(ws, s["list"], f"{L}{T.first}:{L}{T.last}", s.get("warn", False))
    for i, r in enumerate(T.rows()):
        ws.row_dimensions[r].height = row_height
        rowdata = data[i] if i < len(data) else {}
        for s in spec:
            key, kind = s["key"], s["kind"]
            cell = ws[T.c(key, r)]
            fmt = s.get("fmt")
            align = s.get("align")
            if kind == "in":
                v = rowdata.get(key)
                if v is None and "default" in s:
                    v = s["default"]
                cell.value = v
                style(cell, "input", fmt, wrap=s.get("wrap", False), align=align)
            elif kind == "id":
                cell.value = id_start + i
                style(cell, "autogrey", INT, align="center", color=GREY_TEXT)
            elif kind in ("auto", "hid"):
                cell.value = T.f(s["f"], r)
                if kind == "auto":
                    style(cell, "autogrey" if s.get("grey") else "auto", fmt,
                          wrap=s.get("wrap", False), align=align, bold=s.get("bold", False))
                else:
                    style(cell, "hidden", fmt)


# ---------------------------------------------------------------- dropdowns
class Validations:
    """Collect list validations per sheet and flush them."""

    def __init__(self, lists):
        self.lists = lists
        self.pending = {}
        self.force_warn = set()   # lists where typing a value not in the list is allowed

    def add(self, ws, list_name, rng, warn=False):
        warn = warn or list_name in self.force_warn
        key = (ws.title, list_name, warn)
        if key not in self.pending:
            src = self.lists.ref(list_name)
            dv = DataValidation(type="list", formula1=f"={src}", allow_blank=True)
            dv.showErrorMessage = True
            if warn:
                dv.errorStyle = "warning"
                dv.error = "This name is not in the list. Keep it anyway?"
            else:
                dv.error = "Please pick a value from the dropdown list."
            dv.errorTitle = "Renovation OS"
            ws.add_data_validation(dv)
            self.pending[key] = dv
        self.pending[key].add(rng)


class Lists:
    """Registry of dropdown lists on the hidden Lists sheet (one column per list)."""

    def __init__(self):
        self.items = {}   # name -> (column letter, values)
        self.next_col = 1

    def register(self, name, values, header=None):
        L = col_let(self.next_col)
        self.next_col += 1
        self.items[name] = (L, values, header or name)
        return L

    def ref(self, name):
        L, values, _ = self.items[name]
        return f"{q(S_LISTS)}!${L}$2:${L}${1 + len(values)}"

    def col(self, name):
        return self.items[name][0]


# ---------------------------------------------------------------- linter
WHITELIST = {
    "SUM", "SUMIFS", "COUNTIFS", "COUNTIF", "COUNTA", "COUNT", "IF", "IFERROR", "INDEX", "MATCH",
    "SMALL", "LARGE", "MIN", "MAX", "TODAY", "DATE", "YEAR", "MONTH", "DAY", "WEEKDAY", "AND",
    "OR", "NOT", "ROUND", "ABS", "REPT", "LEFT", "LEN", "ROW", "SUMPRODUCT", "FIXED", "HYPERLINK",
}
_STR = re.compile(r'"(?:[^"]|"")*"')
_FN = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
_WHOLE_COL = re.compile(r"(?<![A-Za-z0-9_$])\$?[A-Z]{1,3}:\$?[A-Z]{1,3}(?![0-9A-Za-z_])")


def lint_formula(f):
    body = _STR.sub('""', f)
    problems = []
    for fn in _FN.findall(body):
        if fn.upper() not in WHITELIST:
            problems.append(f"function {fn}")
    if "[" in body or "{" in body:
        problems.append("bracket/array syntax")
    if _WHOLE_COL.search(body):
        problems.append("whole-column reference")
    if "@" in body:
        problems.append("@ operator")
    return problems


_SHEET_REF = re.compile(r"'((?:[^']|'')+)'!")


def lint_workbook(wb):
    errors = []
    sheets = set(wb.sheetnames)
    names = set(wb.defined_names.keys())
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    for p in lint_formula(v):
                        errors.append(f"{ws.title}!{c.coordinate}: {p}: {v[:120]}")
                    for s in _SHEET_REF.findall(v):
                        if s.replace("''", "'") not in sheets:
                            errors.append(f"{ws.title}!{c.coordinate}: missing sheet '{s}'")
    for n, dn in wb.defined_names.items():
        for s in _SHEET_REF.findall(dn.attr_text):
            if s.replace("''", "'") not in sheets:
                errors.append(f"name {n}: missing sheet '{s}'")
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                for fml in rule.formula or []:
                    for p in lint_formula("=" + fml):
                        errors.append(f"{ws.title} CF {cf.sqref}: {p}")
    return errors


# ---------------------------------------------------------------- licence + hardening
AUTHOR = "Hang Nguyen"
COPYRIGHT_YEAR = 2026
LICENSE_LINE = (f"© {COPYRIGHT_YEAR} {AUTHOR} — Renovation OS. Licensed to the original purchaser "
                "for personal use only. Do not share, copy, resell or redistribute.")
LICENSE_TEXT = [
    f"Renovation OS © {COPYRIGHT_YEAR} {AUTHOR}. All rights reserved.",
    "• This file is licensed, not sold, to the original purchaser for personal, non-commercial "
    "use on their own home renovation.",
    "• You may NOT share, copy, resell, redistribute, sublicense, give away or upload this file "
    "(or any modified version or part of it) to any website, marketplace, group or other person.",
    "• The formulas, structure, layout and design are the intellectual property of "
    f"{AUTHOR}. Removing this notice or the sheet protection is not permitted.",
    "• Digital product: all sales are final. No refunds once the file has been delivered.",
    "• Provided “as is” as a planning aid, without warranty. It is not financial, legal or "
    "construction advice — always check figures before you act on them.",
]


def harden(wb, password):
    """Hide + lock every formula, password-protect sheets and workbook structure, stamp licence."""
    from openpyxl.workbook.protection import WorkbookProtection
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and c.protection.locked:
                    c.protection = Protection(locked=True, hidden=True)
        if ws.sheet_state == "visible":
            b3 = ws["B3"]
            assert b3.value is None, f"{ws.title}!B3 is not free for the licence line"
            b3.value = LICENSE_LINE
            b3.font = font(8, False, GREY_TEXT, True)
            b3.protection = LOCKED
            ws.oddFooter.left.text = f"© {COPYRIGHT_YEAR} {AUTHOR} · Renovation OS"
            ws.oddFooter.left.size = 8
            ws.oddFooter.right.text = "Personal licence — do not share or resell"
            ws.oddFooter.right.size = 8
        protect(ws)
        ws.protection.selectLockedCells = True     # locked (formula) cells cannot be selected/copied
        ws.protection.password = password
    wb.security = WorkbookProtection(workbookPassword=password, lockStructure=True)
    pr = wb.properties
    pr.creator = AUTHOR
    pr.lastModifiedBy = AUTHOR
    pr.title = "Renovation OS"
    pr.subject = "Home renovation management workbook"
    pr.description = LICENSE_LINE
    pr.keywords = f"Renovation OS; {AUTHOR}; personal licence"
