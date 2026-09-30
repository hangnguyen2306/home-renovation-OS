"""Execution tabs: Timeline (+Gantt), Rooms, Selections & Orders, Issues & Punch List."""
from openpyxl.styles import Alignment

from . import core as C
from .core import DATE, INT, MONEY, PCT, put, section, style, pri_from_due, write_table
from .layout import (CO, CON, GANTT_FIRST, GANTT_WEEKS, ISS, PAY, ROOM_CARD_H, ROOM_COUNT,
                     ROOMDATA, SEL, TL, TL_ASOF_CELL, room_card)
from .sheets_money import cur_hdr


# ======================================================================= TIMELINE
def build_timeline(ctx):
    ws = ctx.ws(C.S_TL)
    ws.sheet_properties.tabColor = C.TAB_ORANGE
    C.title(ws, "📅 Timeline",
            "What this tab does: tasks with durations and dependencies → start/end dates, "
            "late warnings and a 52-week Gantt chart. Forecast finish feeds the Dashboard.")
    T = TL
    spec = [
        dict(key="id", header="ID", width=5, kind="id",
             note="Fixed task number. Use it in the Predecessor column of later tasks."),
        dict(key="phase", header="Phase", width=15, kind="in", list="Phase"),
        dict(key="task", header="Task", width=30, kind="in"),
        dict(key="room", header="Room", width=14, kind="in", list="Rooms"),
        dict(key="contractor", header="Contractor", width=20, kind="in", list="Contractors",
             warn=True),
        dict(key="dur", header="Duration (days)", width=9, kind="in", fmt=INT, align="center",
             note="Calendar days the task takes."),
        dict(key="pred", header="Predecessor ID", width=10, kind="in", fmt=INT, align="center",
             note="ID of the task that must finish first. It must be a LOWER ID (a task "
                  "listed above this one). Leave empty if none."),
        dict(key="manual", header="Earliest start (optional)", width=12, kind="in", fmt=DATE,
             note="Optional: the task cannot start before this date (e.g. contractor "
                  "availability)."),
        dict(key="start", header="Start", width=12, kind="auto", fmt=DATE, bold=True,
             note="Latest of: project start, your earliest start, predecessor end + 1.",
             f='=IF(OR({task}="",ProjStart=""),"",MAX(ProjStart,IF({manual}<>"",{manual},0),'
               'IF({pred}="",0,IFERROR(INDEX($L$6:L{r1},MATCH({pred},$B$6:B{r1},0))+1,0))))'),
        dict(key="delay", header="Delay (days)", width=8, kind="in", fmt=INT, align="center",
             note="Days this task slipped. Everything that depends on it moves too."),
        dict(key="end", header="End", width=12, kind="auto", fmt=DATE, bold=True,
             f='=IF({start}="","",{start}+MAX(1,{dur})+IF({delay}="",0,{delay})-1)'),
        dict(key="status", header="Status", width=12, kind="in", list="TaskStatus"),
        dict(key="done_date", header="Done date", width=12, kind="in", fmt=DATE),
        dict(key="affects", header="Affects N tasks", width=9, kind="auto", align="center",
             grey=True, note="How many tasks list this one as predecessor. >0 = on the "
                             "critical path: a delay here pushes them.",
             f=f'=IF({{task}}="","",COUNTIF({T.local("pred")},{{id}}))'),
        dict(key="flag", header="Flag", width=16, kind="auto", align="center",
             f='=IF({task}="","",IF(AND({pred}<>"",{pred}>={id}),"⚠️ Predecessor must be above",'
               'IF({status}="Done","✓ Done",IF({code}=3,"🔴 Late",IF({status}="In progress",'
               '"In progress","")))))'),
        dict(key="code", header="code", kind="hid",
             f='=IF(OR({task}="",{end}=""),"",IF({status}="Done",2,IF(AsOf>{end},3,1)))'),
        dict(key="first", header="first", kind="hid",
             f='=IF(OR({room}="",{contractor}=""),0,IF(COUNTIFS($E$7:E{r},{room},$F$7:F{r},'
               '{contractor})=1,1,0))'),
        dict(key="okey", header="okey", kind="hid",
             f='=IF({first}=1,{room}&"|"&SUMIFS($R$7:R{r},$E$7:E{r},{room}),"")'),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f='=IF({code}="","",IF({code}=3,{end},IF(OR({status}="Not started",{status}=""),'
               '{start},"")))'),
        dict(key="pri", header="pri", kind="hid",
             f=f'=IF({{due}}="","",IF({{code}}=3,1,{pri_from_due("{due}")}))'),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","",IF({code}=3,"Late task: "&{task},IF({due}<AsOf,'
               '"Should have started: ","Starts soon: ")&{task})&IF({contractor}="",""," — "&'
               '{contractor}))'),
    ]
    data = ctx.data.TASKS if ctx.demo else ctx.data.BLANK_TASKS
    write_table(ws, T, spec, data, ctx.dv)
    ws[TL_ASOF_CELL] = "=AsOf"
    ws[TL_ASOF_CELL].number_format = DATE

    # summary strip
    Lend = T.L("end")
    put(ws, "C4", "Forecast finish", "label", bold=True, align="right")
    put(ws, "D4", f'=IF(COUNT({Lend}{T.first}:{Lend}{T.last})=0,"",MAX({Lend}{T.first}:{Lend}{T.last}))',
        "band", DATE, bold=True, size=12, align="center")
    C.define(ctx.wb, "ForecastEnd", C.S_TL, "$D$4")
    put(ws, "E4", "Target move-in", "label", bold=True, align="right")
    ws.merge_cells("E4:F4")
    put(ws, "G4", "=IF(TargetEnd=\"\",\"\",TargetEnd)", "autogrey", DATE, bold=True)
    ws.merge_cells("G4:H4")
    put(ws, "I4", '=IF(OR(ForecastEnd="",TargetEnd=""),"",IF(ForecastEnd>TargetEnd,"🔴 "&'
        '(ForecastEnd-TargetEnd)&" days late",IF(ForecastEnd=TargetEnd,"🟢 on target","🟢 "&'
        '(TargetEnd-ForecastEnd)&" days early")))', "label", bold=True)
    ws.merge_cells("I4:L4")
    ws.row_dimensions[4].height = 26

    # Gantt
    g0 = GANTT_FIRST
    gl = C.col_add(g0, GANTT_WEEKS - 1)
    ws.column_dimensions[C.col_add(g0, -1)].width = 2
    put(ws, f"{g0}4", "Gantt: teal = planned · grey = done · red = late · yellow column = "
        "this week", "note")
    for i in range(GANTT_WEEKS):
        L = C.col_add(g0, i)
        prev = C.col_add(g0, i - 1)
        ws.column_dimensions[L].width = 3.3
        h = ws[f"{L}6"]
        if i == 0:
            h.value = ('=IF(ProjStart="","",ProjStart-IF(WeekStart="Sunday",WEEKDAY(ProjStart,1)-1,'
                       'WEEKDAY(ProjStart,3)))')
        else:
            h.value = f'=IF({prev}6="","",{prev}6+7)'
        style(h, "header", "dd mmm")
        h.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        for r in T.rows():
            ws[f"{L}{r}"].border = C.Border(left=C.Side(style="hair", color="E5E7EB"))
    ws.row_dimensions[6].height = 48
    rng = f"{g0}{T.first}:{gl}{T.last}"
    ov =(f"AND({g0}$6<>\"\",$J{T.first}<>\"\",$J{T.first}<={g0}$6+6,$L{T.first}>={g0}$6)")
    C.add_cf_fill(ws, rng, f"AND({ov},$Q{T.first}=3)", "DC2626", stop=True)
    C.add_cf_fill(ws, rng, f"AND({ov},$Q{T.first}=2)", "9CA3AF", stop=True)
    C.add_cf_fill(ws, rng, ov, "14B8A6", stop=True)
    today = f"${TL_ASOF_CELL[0]}${TL_ASOF_CELL[1:]}"
    C.add_cf_fill(ws, rng, f'AND({g0}$6<>"",{today}>={g0}$6,{today}<{g0}$6+7)', "FDE68A")
    C.add_cf_fill(ws, f"{g0}6:{gl}6", f'AND({g0}$6<>"",{today}>={g0}$6,{today}<{g0}$6+7)',
                  "F59E0B", C.WHITE)
    fl = T.L("flag")
    frng = f"{fl}{T.first}:{fl}{T.last}"
    C.add_cf(ws, frng, f"$Q{T.first}=3", C.BAD, bold=True)
    C.add_cf(ws, frng, f"$Q{T.first}=2", C.GOOD)
    C.add_cf(ws, frng, f'LEFT({fl}{T.first},1)="⚠"', C.BAD, bold=True)
    ws.freeze_panes = f"E{T.first}"
    C.protect(ws)


# ======================================================================= ROOMS
def build_rooms(ctx):
    ws = ctx.ws(C.S_ROOMS)
    ws.sheet_properties.tabColor = C.TAB_ORANGE
    ws.sheet_view.showGridLines = False
    C.title(ws, "🛋️ Rooms",
            "What this tab does: one card per room — money, progress, decisions and problems. "
            "Rename rooms in the yellow title cells; every dropdown follows.")
    C.widths(ws, {"B": 15, "C": 11, "D": 13, "E": 15, "F": 13, "G": 3, "H": 15, "I": 11,
                  "J": 13, "K": 15, "L": 13, "M": 3})
    D = ROOMDATA
    rooms = ctx.data.ROOMS if ctx.demo else [(n, None, None) for n in ctx.data.BLANK_ROOM_NAMES]
    bar = lambda x: f'REPT("█",ROUND(MIN(1,{x})*10,0))&REPT("░",10-ROUND(MIN(1,{x})*10,0))'  # noqa
    for k in range(ROOM_COUNT):
        c0, c1, t = room_card(k)
        cc = [C.col_add(c0, i) for i in range(5)]
        r = D.first + k
        dref = lambda key: D.c(key, r)  # noqa: E731
        name, area, budget = rooms[k]
        # card background
        for rr in range(t, t + ROOM_CARD_H):
            for L in cc:
                cell = ws[f"{L}{rr}"]
                cell.fill = C.fill("FAFAFA")
                cell.protection = C.LOCKED
            ws.row_dimensions[rr].height = 21
        for L in cc:
            ws[f"{L}{t}"].border = C.Border(top=C.Side(style="thick", color=C.ACCENT))
            ws[f"{L}{t + ROOM_CARD_H - 1}"].border = C.Border(bottom=C.Side(style="thin", color=C.LINE))
        # title
        ws.merge_cells(f"{cc[0]}{t}:{cc[3]}{t}")
        put(ws, f"{cc[0]}{t}", name, "input", size=13, bold=True, border=False)
        ws.row_dimensions[t].height = 28
        put(ws, f"{cc[4]}{t}", f'=IF({dref("lvl")}="","",IF({dref("lvl")}=0,"🟢 OK",'
            f'IF({dref("lvl")}=1,"🟡 Watch","🔴 Act")))', "auto", align="center", bold=True)
        C.level_cf(ws, f"{cc[4]}{t}", f"${D.L('lvl')}${r}")

        def lab(rr, col, text):
            c = ws[f"{col}{rr}"]
            c.value = text
            c.font = C.font(9, False, C.GREY_TEXT)
            c.alignment = Alignment(vertical="center", indent=1)

        def val(rr, col, f, fmt=None, bold=True, kind="label"):
            c = ws[f"{col}{rr}"]
            c.value = f
            c.font = C.font(11, bold, C.DARK)
            c.alignment = Alignment(horizontal="right", vertical="center")
            if fmt:
                c.number_format = fmt

        lab(t + 1, cc[0], '="Floor area ("&AreaUnit&")"')
        put(ws, f"{cc[2]}{t + 1}", area, "input", C.NUM)
        lab(t + 1, cc[3], '="Budget ("&CurSym&")"')
        put(ws, f"{cc[4]}{t + 1}", budget, "input", MONEY)
        lab(t + 2, cc[0], '="Committed ("&CurSym&")"')
        val(t + 2, cc[2], f"={dref('committed')}", MONEY)
        lab(t + 2, cc[3], '="Paid ("&CurSym&")"')
        val(t + 2, cc[4], f"={dref('paid')}", MONEY)
        lab(t + 3, cc[0], "Budget used")
        ws.merge_cells(f"{cc[2]}{t + 3}:{cc[3]}{t + 3}")
        b = ws[f"{cc[2]}{t + 3}"]
        pct_b = f'IF(OR({dref("budget")}="",{dref("budget")}=0),0,{dref("committed")}/{dref("budget")})'
        b.value = f'=IF(OR({dref("budget")}="",{dref("budget")}=0),"",{bar(pct_b)})'
        b.font = C.font(11, False, C.ACCENT)
        val(t + 3, cc[4], f'=IF(OR({dref("budget")}="",{dref("budget")}=0),"–",{pct_b})', PCT)
        C.add_cf(ws, f"{cc[2]}{t + 3}:{cc[4]}{t + 3}", f"{pct_b.replace(dref('budget'), '$' + D.L('budget') + '$' + str(r)).replace(dref('committed'), '$' + D.L('committed') + '$' + str(r))}>1", C.BAD, bold=True)
        lab(t + 4, cc[0], "Progress")
        ws.merge_cells(f"{cc[2]}{t + 4}:{cc[3]}{t + 4}")
        p = ws[f"{cc[2]}{t + 4}"]
        p.value = f'=IF({dref("progress")}="","",{bar(dref("progress"))})'
        p.font = C.font(11, False, "14B8A6")
        val(t + 4, cc[4], f'=IF({dref("progress")}="","–",{dref("progress")})', PCT)
        lab(t + 5, cc[0], "Tasks done")
        val(t + 5, cc[2], f'={dref("t_done")}&" / "&{dref("t_total")}')
        lab(t + 5, cc[3], "Late tasks")
        val(t + 5, cc[4], f"={dref('late')}", INT)
        lab(t + 6, cc[0], "Decisions waiting")
        val(t + 6, cc[2], f"={dref('decisions')}", INT)
        lab(t + 6, cc[3], "Open issues")
        val(t + 6, cc[4], f"={dref('open_issues')}", INT)
        lab(t + 7, cc[0], '="Cost per "&AreaUnit')
        val(t + 7, cc[2], f"={dref('per_area')}", MONEY)
        lab(t + 7, cc[3], "High-severity")
        val(t + 7, cc[4], f"={dref('open_high')}", INT)
        lab(t + 8, cc[0], "Contractors")
        ws.merge_cells(f"{cc[1]}{t + 8}:{cc[4]}{t + 8}")
        cn = ws[f"{cc[1]}{t + 8}"]
        cn.value = (f'=IF({dref("c1")}="","–",{dref("c1")}&IF({dref("c2")}="","",", "&{dref("c2")})'
                    f'&IF({dref("c3")}="","",", "&{dref("c3")}))')
        cn.font = C.font(9, False, C.DARK)
        cn.alignment = Alignment(horizontal="right", vertical="center", shrink_to_fit=True)
        for rr, col in ((t + 5, cc[4]), (t + 7, cc[4])):
            C.add_cf(ws, f"{col}{rr}", f"{col}{rr}>0", C.BAD, bold=True)
        C.add_cf(ws, f"{cc[4]}{t + 6}", f"{cc[4]}{t + 6}>0", C.WARN, bold=True)
        C.add_cf(ws, f"{cc[2]}{t + 6}", f"{cc[2]}{t + 6}>0", C.WARN, bold=True)

        # hidden data row
        nm = dref("name")
        ws[nm] = f'=IF({cc[0]}{t}="","",{cc[0]}{t})'
        ws[dref("area")] = f'=IF({cc[2]}{t + 1}="","",{cc[2]}{t + 1})'
        ws[dref("budget")] = f'=IF({cc[4]}{t + 1}="","",{cc[4]}{t + 1})'
        ws[dref("committed")] = (
            f'=IF({nm}="","",SUMIFS({CON.a("value")},{CON.a("room")},{nm},{CON.a("signed")},"Yes")'
            f'+SUMIFS({SEL.a("price")},{SEL.a("room")},{nm})+SUMIFS({CO.a("cost")},{CO.a("room")},'
            f'{nm},{CO.a("status")},"Approved"))')
        ws[dref("paid")] = f'=IF({nm}="","",SUMIFS({PAY.a("paid_amt")},{PAY.a("room_auto")},{nm}))'
        ws[dref("t_total")] = f'=IF({nm}="",0,COUNTIFS({TL.a("room")},{nm}))'
        ws[dref("t_done")] = f'=IF({nm}="",0,COUNTIFS({TL.a("room")},{nm},{TL.a("status")},"Done"))'
        ws[dref("progress")] = f'=IF({dref("t_total")}=0,"",{dref("t_done")}/{dref("t_total")})'
        ws[dref("decisions")] = f'=IF({nm}="",0,SUMIFS({SEL.a("dec_wait")},{SEL.a("room")},{nm}))'
        ws[dref("open_issues")] = f'=IF({nm}="",0,SUMIFS({ISS.a("open")},{ISS.a("room")},{nm}))'
        ws[dref("open_high")] = f'=IF({nm}="",0,SUMIFS({ISS.a("open_high")},{ISS.a("room")},{nm}))'
        ws[dref("late")] = f'=IF({nm}="",0,COUNTIFS({TL.a("room")},{nm},{TL.a("code")},3))'
        bud = dref("budget")
        ws[dref("lvl")] = (f'=IF({nm}="","",IF(OR({dref("open_high")}>0,{dref("late")}>0,'
                           f'AND({bud}<>"",{dref("committed")}>IF({bud}="",0,{bud}),'
                           f'IF({bud}="",0,{bud})>0)),2,IF(OR({dref("open_issues")}>0,'
                           f'{dref("decisions")}>0),1,0)))')
        for i, key in enumerate(("c1", "c2", "c3")):
            ws[dref(key)] = (f'=IF({nm}="","",IFERROR(INDEX({TL.a("contractor")},MATCH({nm}&"|{i + 1}",'
                             f'{TL.a("okey")},0)),""))')
        ws[dref("per_area")] = (f'=IF(OR({nm}="",{dref("area")}=""),"",IF({dref("area")}=0,"",'
                                f'{dref("committed")}/{dref("area")}))')
        ws[dref("forecast_room")] = f'=IF({nm}="","",MAX(IF({bud}="",0,{bud}),{dref("committed")}))'
    for key, L in D.cols.items():
        ws[f"{L}{D.header_row}"] = key
        ws.column_dimensions[L].hidden = True
    C.protect(ws)


# ======================================================================= SELECTIONS
def build_selections(ctx):
    ws = ctx.ws(C.S_SEL)
    ws.sheet_properties.tabColor = C.TAB_ORANGE
    C.title(ws, "🧩 Selections & Orders",
            "What this tab does: every choice you must make (tiles, taps, appliances…) with "
            "automatic decide-by and order-by dates so nothing holds up the builders.")
    T = SEL
    st = {"DL": C.ST_DELIV_LATE, "LATE": C.ST_ORDER_LATE}
    status_f = (
        '=IF({item}="","",IF({delivered}="Yes","%s",IF({ordered}="Yes",IF(AND({delivery}<>"",'
        '{need}<>""),IF({delivery}>{need},"%s","%s"),"%s"),IF({selected}="",IF({decide_by}="",'
        '"%s",IF({decide_by}<=AsOf+AlertDays,"%s","%s")),IF({order_by}="","%s",IF(AsOf>{order_by},'
        '"%s",IF({order_by}<=AsOf+AlertDays,"%s","%s")))))))'
        % (C.ST_DELIVERED, C.ST_DELIV_LATE, C.ST_ORDERED, C.ST_ORDERED, C.ST_WAITING,
           C.ST_DECIDE, C.ST_WAITING, C.ST_SELECTED, C.ST_ORDER_LATE, C.ST_ORDER_NOW,
           C.ST_SELECTED))
    spec = [
        dict(key="item", header="Item / decision", width=26, kind="in"),
        dict(key="room", header="Room", width=13, kind="in", list="Rooms"),
        dict(key="cat", header="Budget category", width=14, kind="in", list="Categories"),
        dict(key="a_name", header="Option A", width=15, kind="in"),
        dict(key="a_price", header=cur_hdr("A"), width=9, kind="in", fmt=MONEY),
        dict(key="b_name", header="Option B", width=15, kind="in"),
        dict(key="b_price", header=cur_hdr("B"), width=9, kind="in", fmt=MONEY),
        dict(key="c_name", header="Option C", width=15, kind="in"),
        dict(key="c_price", header=cur_hdr("C"), width=9, kind="in", fmt=MONEY),
        dict(key="selected", header="Chosen", width=8, kind="in", list="ABC", align="center",
             note="Pick A, B or C once decided. Empty = decision still open."),
        dict(key="price", header=cur_hdr("Chosen price"), width=11, kind="auto", fmt=MONEY,
             grey=True, note="Counts as Materials on the Budget tab.",
             f='=IF({selected}="A",{a_price},IF({selected}="B",{b_price},IF({selected}="C",{c_price},"")))'),
        dict(key="reason", header="Why this option", width=20, kind="in"),
        dict(key="supplier", header="Supplier", width=16, kind="in"),
        dict(key="lead", header="Lead time (days)", width=9, kind="in", fmt=INT, align="center",
             note="Days between ordering and delivery (ask the supplier)."),
        dict(key="task", header="Timeline task ID", width=9, kind="in", fmt=INT, align="center",
             note="ID of the Timeline task that installs this item. Its start date becomes the "
                  "needed-on date."),
        dict(key="need_in", header="Needed-on (override)", width=12, kind="in", fmt=DATE,
             note="Optional: type a date to override the Timeline date."),
        dict(key="need", header="Needed on", width=12, kind="auto", fmt=DATE, grey=True,
             f=f'=IF({{item}}="","",IF({{need_in}}<>"",{{need_in}},IF({{task}}="","",IFERROR('
               f'INDEX({TL.a("start")},MATCH({{task}},{TL.a("id")},0)),""))))'),
        dict(key="order_by", header="ORDER BY", width=12, kind="auto", fmt=DATE, bold=True,
             note="Needed-on − lead time − order safety buffer (Start Here).",
             f='=IF({need}="","",{need}-IF({lead}="",0,{lead})-OrderBuf)'),
        dict(key="decide_by", header="DECIDE BY", width=12, kind="auto", fmt=DATE, bold=True,
             note="Order-by − decide-before-order buffer (Start Here).",
             f='=IF({order_by}="","",{order_by}-DecideBuf)'),
        dict(key="ordered", header="Ordered?", width=9, kind="in", list="YesNo", align="center"),
        dict(key="order_date", header="Order date", width=12, kind="in", fmt=DATE),
        dict(key="delivery", header="Expected delivery", width=12, kind="in", fmt=DATE),
        dict(key="delivered", header="Delivered?", width=9, kind="in", list="YesNo",
             align="center"),
        dict(key="status", header="Status", width=28, kind="auto", bold=True, f=status_f),
        dict(key="lvl", header="lvl", kind="hid",
             f=f'=IF({{item}}="","",IF(OR({{late_mat}}=1,{{dec_over}}=1),2,IF(OR({{order_now}}=1,'
               f'{{dec_soon}}=1),1,IF(OR({{status}}="{C.ST_ORDERED}",{{status}}="{C.ST_DELIVERED}"),0,""))))'),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f=f'=IF({{item}}="","",IF({{late_mat}}=1,IF({{status}}="{st["DL"]}",{{need}},{{order_by}}),'
               f'IF({{dec_wait}}=1,{{decide_by}},IF({{need_order}}=1,{{order_by}},""))))'),
        dict(key="pri", header="pri", kind="hid",
             f=f'=IF({{due}}="","",IF({{late_mat}}=1,1,{pri_from_due("{due}")}))'),
        dict(key="text", header="text", kind="hid",
             f=f'=IF({{pri}}="","",IF({{late_mat}}=1,IF({{status}}="{st["DL"]}","Delivery after '
               f'install date: "&{{item}},"Order late — install at risk: "&{{item}}),IF({{dec_wait}}=1,'
               f'IF({{pri}}=1,"Decision overdue: ","Decide: ")&{{item}},"Order: "&{{item}}&'
               f'IF({{supplier}}="",""," from "&{{supplier}}))))'),
        dict(key="dec_wait", header="dec_wait", kind="hid",
             f='=IF(AND({item}<>"",{selected}="",{ordered}<>"Yes",{delivered}<>"Yes"),1,0)'),
        dict(key="dec_over", header="dec_over", kind="hid",
             f='=IF(AND({dec_wait}=1,{decide_by}<>""),IF({decide_by}<AsOf,1,0),0)'),
        dict(key="dec_soon", header="dec_soon", kind="hid",
             f='=IF(AND({dec_wait}=1,{decide_by}<>""),IF(AND({decide_by}>=AsOf,'
               '{decide_by}<=AsOf+AlertDays),1,0),0)'),
        dict(key="order_now", header="order_now", kind="hid",
             f=f'=IF({{status}}="{C.ST_ORDER_NOW}",1,0)'),
        dict(key="late_mat", header="late_mat", kind="hid",
             f=f'=IF(OR({{status}}="{st["LATE"]}",{{status}}="{st["DL"]}"),1,0)'),
        dict(key="need_order", header="need_order", kind="hid",
             f='=IF(AND({item}<>"",{selected}<>"",{ordered}<>"Yes",{delivered}<>"Yes"),1,0)'),
        dict(key="order_late", header="order_late", kind="hid",
             f=f'=IF({{status}}="{st["LATE"]}",1,0)'),
    ]
    write_table(ws, T, spec, ctx.data.SELECTIONS if ctx.demo else [], ctx.dv)
    S = T.L("status")
    C.level_cf(ws, f"{S}{T.first}:{S}{T.last}", f"${T.L('lvl')}{T.first}")
    for key in ("order_by", "decide_by"):
        L = T.L(key)
        C.add_cf(ws, f"{L}{T.first}:{L}{T.last}", f'AND({L}{T.first}<>"",${T.L("lvl")}{T.first}=2)',
                 C.BAD, bold=True)

    # counters
    cnt = [
        ("B", "C", "Decisions waiting", f"=SUM({T.local('dec_wait')})", "DecisionsWaiting"),
        ("E", "G", '="Orders due in next "&AlertDays&" days"', f"=SUM({T.local('order_now')})",
         "OrdersDue"),
        ("J", "M", "Late orders — install at risk", f"=SUM({T.local('order_late')})",
         "OrdersLate"),
        ("P", "R", "Deliveries after install date",
         f"=SUM({T.local('late_mat')})-SUM({T.local('order_late')})", "DeliveriesLate"),
    ]
    for c0, c1, label, f, name in cnt:
        ws.merge_cells(f"{c0}4:{c1}4")
        put(ws, f"{c0}4", label, "label", color=C.GREY_TEXT)
        ws.merge_cells(f"{c0}5:{c1}5")
        c = put(ws, f"{c0}5", f, "band", INT, size=16, bold=True, align="left")
        C.define(ctx.wb, name, C.S_SEL, f"${c0}$5")
        C.add_cf(ws, f"{c0}5", f"{c0}5>0", C.BAD if name in ("OrdersLate", "DeliveriesLate")
                 else C.WARN, bold=True)
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = f"C{T.first}"
    C.protect(ws)


# ======================================================================= ISSUES
def build_issues(ctx):
    ws = ctx.ws(C.S_ISS)
    ws.sheet_properties.tabColor = C.TAB_ORANGE
    C.title(ws, "🛠️ Issues & Punch List",
            "What this tab does: every problem found on site and every snag for the final "
            "punch list — with who fixes it, by when, and what it may cost.")
    T = ISS
    spec = [
        dict(key="no", header="#", width=5, kind="id"),
        dict(key="type", header="Type", width=11, kind="in", list="IssueType"),
        dict(key="found", header="Date found", width=12, kind="in", fmt=DATE),
        dict(key="room", header="Room", width=14, kind="in", list="Rooms"),
        dict(key="desc", header="Description", width=34, kind="in", wrap=True),
        dict(key="resp", header="Responsible", width=20, kind="in", list="Contractors",
             warn=True),
        dict(key="sev", header="Severity", width=10, kind="in", list="Severity", align="center",
             note="High = safety, structure, water, or blocks other work. Always shows as 🔴."),
        dict(key="cost", header=cur_hdr("Estimated cost"), width=12, kind="in", fmt=MONEY,
             note="What it may cost YOU (0 if the contractor pays). Open issue costs are shown "
                  "as risk on the Budget tab."),
        dict(key="action", header="Next action", width=26, kind="in", wrap=True),
        dict(key="act_date", header="Action date", width=12, kind="in", fmt=DATE),
        dict(key="status", header="Status", width=12, kind="in", list="IssueStatus"),
        dict(key="res_date", header="Resolved date", width=12, kind="in", fmt=DATE),
        dict(key="link", header="Photo link", width=22, kind="in",
             note="Paste a Google Photos / Drive / OneDrive link."),
        dict(key="open", header="open", kind="hid",
             f='=IF(AND({desc}<>"",{status}<>"Resolved"),1,0)'),
        dict(key="lvl", header="lvl", kind="hid",
             f='=IF({desc}="","",IF({open}=0,0,IF({sev}="High",2,1)))'),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f='=IF({open}=1,IF({act_date}<>"",{act_date},IF({sev}="High",AsOf,"")),"")'),
        dict(key="pri", header="pri", kind="hid",
             f=f'=IF({{due}}="","",IF({{sev}}="High",1,{pri_from_due("{due}")}))'),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","",IF({sev}="High","High-severity issue: ",IF({type}="Punch item",'
               '"Punch item: ","Issue action due: "))&{desc}&IF({action}="",""," → "&{action}))'),
        dict(key="open_cost", header="open_cost", kind="hid",
             f='=IF({open}=1,IF({cost}="",0,{cost}),0)'),
        dict(key="open_high", header="open_high", kind="hid",
             f='=IF(AND({open}=1,{sev}="High"),1,0)'),
        dict(key="open_punch", header="open_punch", kind="hid",
             f='=IF(AND({open}=1,{type}="Punch item"),1,0)'),
    ]
    write_table(ws, T, spec, ctx.data.ISSUES if ctx.demo else [], ctx.dv, row_height=24)
    S = T.L("status")
    C.level_cf(ws, f"{S}{T.first}:{S}{T.last}", f"${T.L('lvl')}{T.first}")
    Sv = T.L("sev")
    C.add_cf(ws, f"{Sv}{T.first}:{Sv}{T.last}", f'AND({Sv}{T.first}="High",${T.L("open")}{T.first}=1)',
             C.BAD, bold=True)
    cnt = [
        ("B", "C", "Open — High", f'=COUNTIFS({T.local("open")},1,{T.local("sev")},"High")',
         "OpenHigh", C.BAD),
        ("D", "E", "Open — Medium", f'=COUNTIFS({T.local("open")},1,{T.local("sev")},"Medium")',
         None, C.WARN),
        ("F", "F", "Open — Low", f'=COUNTIFS({T.local("open")},1,{T.local("sev")},"Low")', None,
         C.WARN),
        ("G", "G", "Open punch items", f"=SUM({T.local('open_punch')})", None, C.WARN),
        ("I", "J", '="Est. cost of open issues ("&CurSym&")"', "=IssuesOpenCost", None, C.WARN),
    ]
    for c0, c1, label, f, name, colr in cnt:
        if c0 != c1:
            ws.merge_cells(f"{c0}4:{c1}4")
            ws.merge_cells(f"{c0}5:{c1}5")
        put(ws, f"{c0}4", label, "label", color=C.GREY_TEXT)
        put(ws, f"{c0}5", f, "band", MONEY if c0 == "I" else INT, size=16, bold=True,
            align="left")
        C.add_cf(ws, f"{c0}5", f"{c0}5>0", colr, bold=True)
        if name:
            C.define(ctx.wb, name, C.S_ISS, f"${c0}$5")
    C.define(ctx.wb, "OpenIssues", C.S_ISS, f"${T.L('open')}${T.first}:${T.L('open')}${T.last}")
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = f"D{T.first}"
    C.protect(ws)
