"""Money tabs: Budget, Quotes, Contractors, Payments, Change Orders."""
from openpyxl.styles import Alignment

from . import core as C
from .core import MONEY, DATE, PCT, INT, put, section, style, pri_from_due, write_table
from .layout import (BUDGET, BUDGET_TOTAL_ROW, BUDGET_ALERT_ROWS, CASH, CASH30_CELL, CHECKLIST,
                     CO, COMM, CON, ISS, PAY, PAYSUM, QO, QSUM_FIRST, QEXP_FIRST, QUOTE_BLOCKS,
                     QUOTE_VCOLS, QUOTE_XCOLS, SEL, qblock)


def cur_hdr(text):
    return f'="{text} ("&CurSym&")"'


# ======================================================================= BUDGET
def build_budget(ctx):
    ws = ctx.ws(C.S_BUDGET)
    ws.sheet_properties.tabColor = C.TAB_GREEN
    C.title(ws, "💶 Budget",
            "What this tab does: compares your budget per category with what you have "
            "committed, invoiced and paid, and forecasts the final cost.")
    T = BUDGET
    tot = BUDGET_TOTAL_ROW
    spec = [
        dict(key="cat", header="Category", width=22, kind="in",
             note="Rename, add or clear categories (up to 30). They feed every Category "
                  "dropdown in the workbook."),
        dict(key="budget", header=cur_hdr("Budget"), width=13, kind="in", fmt=MONEY),
        dict(key="stage", header="Stage", width=13, kind="in", list="Stage", align="center",
             note="Estimating = still collecting quotes. Contracted = contract signed. "
                  "Complete = work finished and final invoice received."),
        dict(key="best", header=cur_hdr("Best quote"), width=13, kind="auto", fmt=MONEY,
             grey=True, note="Lowest TRUE cost (incl. hidden costs) from the Quotes tab for "
                             "this category.",
             f='=IF({cat}="","",IFERROR(1/SUMPRODUCT(MAX((QBlockCat={cat})*QBlockInv)),""))'),
        dict(key="contracted", header=cur_hdr("Contracted"), width=13, kind="auto",
             fmt=MONEY, grey=True, note="Sum of SIGNED contracts on the Contractors tab with "
                                        "this category.",
             f=f'=IF({{cat}}="","",SUMIFS({CON.a("value")},{CON.a("cat")},{{cat}},'
               f'{CON.a("signed")},"Yes"))'),
        dict(key="materials", header=cur_hdr("Materials"), width=13, kind="auto", fmt=MONEY,
             grey=True, note="Materials you buy yourself: selected option prices on "
                             "Selections & Orders.",
             f=f'=IF({{cat}}="","",SUMIFS({SEL.a("price")},{SEL.a("cat")},{{cat}}))'),
        dict(key="cos", header=cur_hdr("Approved changes"), width=13, kind="auto", fmt=MONEY,
             grey=True, note="Approved change orders for this category.",
             f=f'=IF({{cat}}="","",SUMIFS({CO.a("cost")},{CO.a("cat")},{{cat}},'
               f'{CO.a("status")},"Approved"))'),
        dict(key="committed", header=cur_hdr("Committed"), width=13, kind="auto", fmt=MONEY,
             bold=True, note="Contracted + materials + approved changes: money you have "
                             "promised to spend.",
             f='=IF({cat}="","",{contracted}+{materials}+{cos})'),
        dict(key="invoiced", header=cur_hdr("Invoiced"), width=13, kind="auto", fmt=MONEY,
             grey=True, f=f'=IF({{cat}}="","",SUMIFS({PAY.a("inv_amt")},{PAY.a("cat")},{{cat}}))'),
        dict(key="paid", header=cur_hdr("Paid"), width=13, kind="auto", fmt=MONEY, grey=True,
             f=f'=IF({{cat}}="","",SUMIFS({PAY.a("paid_amt")},{PAY.a("cat")},{{cat}}))'),
        dict(key="forecast", header=cur_hdr("Forecast"), width=13, kind="auto", fmt=MONEY,
             bold=True,
             note="Estimating: the highest of budget, best quote, committed and invoiced "
                  "(be pessimistic until signed). Contracted/Complete: committed (or invoiced "
                  "if higher).",
             f='=IF({cat}="","",IF(OR({stage}="Estimating",{stage}=""),'
               'MAX({budget},{best},{committed},{invoiced}),MAX({committed},{invoiced})))'),
        dict(key="variance", header=cur_hdr("Variance"), width=13, kind="auto", fmt=MONEY,
             note="Budget − forecast. Negative (red) = over budget.",
             f='=IF({cat}="","",{budget}-{forecast})'),
        dict(key="status", header="Status", width=14, kind="auto", align="center",
             f='=IF({lvl}="","",IF({lvl}=0,"🟢 On track",IF({lvl}=1,"🟡 Watch","🔴 Over")))'),
        dict(key="lvl", header="lvl", kind="hid",
             f='=IF({cat}="","",IF({forecast}<={budget},0,'
               'IF({forecast}<={budget}*(1+BudgetWarnPct),1,2)))'),
        dict(key="chartkey", header="chart key", kind="hid",
             f='=IF({cat}="","",IF(MAX({budget},{forecast})>0,MAX({budget},{forecast})+ROW()/100000,""))'),
    ]
    data = []
    for i, (name, bud, stage) in enumerate(ctx.data.CATEGORIES):
        row = {"cat": name}
        if ctx.demo:
            row.update(budget=bud, stage=stage)
        else:
            row.update(stage="Estimating")
        data.append(row)
    write_table(ws, T, spec, data, ctx.dv)
    for L in ("due", "pri", "text"):
        ws.column_dimensions[T.L(L)].hidden = True

    # totals
    ws.row_dimensions[tot].height = 26
    put(ws, f"B{tot}", "TOTAL", "band", bold=True)
    for key in ("budget", "contracted", "materials", "cos", "committed", "invoiced", "paid",
                "forecast", "variance"):
        L = T.L(key)
        put(ws, f"{L}{tot}", f"=SUM({L}{T.first}:{L}{T.last})", "band", MONEY, bold=True)
    for key in ("stage", "best"):
        put(ws, f"{T.L(key)}{tot}", None, "band")
    put(ws, f"{T.L('lvl')}{tot}",
        f'=IF({T.L("forecast")}{tot}<={T.L("budget")}{tot},0,IF({T.L("forecast")}{tot}<='
        f'{T.L("budget")}{tot}*(1+BudgetWarnPct),1,2))', "hidden")
    put(ws, f"{T.L('status')}{tot}",
        f'=IF({T.L("lvl")}{tot}=0,"🟢 On track",IF({T.L("lvl")}{tot}=1,"🟡 Watch","🔴 Over"))',
        "band", align="center", bold=True)
    for k in ("budget", "committed", "invoiced", "paid", "forecast"):
        pass
    C.define(ctx.wb, "BudgetSum", C.S_BUDGET, f"${T.L('budget')}${tot}")
    C.define(ctx.wb, "CommittedTotal", C.S_BUDGET, f"${T.L('committed')}${tot}")
    C.define(ctx.wb, "InvoicedTotal", C.S_BUDGET, f"${T.L('invoiced')}${tot}")
    C.define(ctx.wb, "PaidTotal", C.S_BUDGET, f"${T.L('paid')}${tot}")
    C.define(ctx.wb, "ForecastTotal", C.S_BUDGET, f"${T.L('forecast')}${tot}")

    # risk block
    F = T.L("forecast")
    risk = [
        (40, "Pending change orders (not approved yet)", "PendingCO",
         f'=SUMIFS({CO.a("cost")},{CO.a("status")},"Pending")', MONEY),
        (41, "Estimated cost of open issues", "IssuesOpenCost",
         f'=SUM({ISS.a("open_cost")})', MONEY),
        (42, "RISK (could still be added)", "RiskTotal", f"={F}40+{F}41", MONEY),
        (43, "FORECAST INCL. RISK", "ForecastRisk", f"=ForecastTotal+{F}42", MONEY),
        (44, "Contingency used by forecast", "ContUsed",
         "=IF(ContAmt>0,MAX(0,ForecastTotal-WorkBudget)/ContAmt,0)", PCT),
    ]
    section(ws, "B39", "RISK — what could still come on top of the forecast", "N")
    for row, label, name, f, fmt in risk:
        ws.merge_cells(f"B{row}:{C.col_add(F, -1)}{row}")
        put(ws, f"B{row}", label, "label", bold=row in (42, 43), align="right")
        put(ws, f"{F}{row}", f, "autogrey" if row < 42 else "band", fmt, bold=True)
        C.define(ctx.wb, name, C.S_BUDGET, f"${F}${row}")
        ws.row_dimensions[row].height = 22
    put(ws, f"{T.L('variance')}44", '=REPT("█",ROUND(MIN(1,ContUsed)*10,0))&REPT("░",10-ROUND(MIN(1,ContUsed)*10,0))',
        "label", color=C.ACCENT)
    C.add_cf(ws, f"{F}44", f"{F}44>0.75", C.BAD, bold=True)

    # budget alerts (hidden helper cells)
    r1, r2 = BUDGET_ALERT_ROWS
    d, p, t = T.L("due"), T.L("pri"), T.L("text")
    ws[f"{d}{r1}"] = '=IF(AND(WorkBudget>0,ForecastTotal>WorkBudget),AsOf,"")'
    ws[f"{p}{r1}"] = f'=IF({d}{r1}="","",1)'
    ws[f"{t}{r1}"] = (f'=IF({p}{r1}="","","Forecast is over your working budget by "&CurSym&" "'
                      f'&FIXED(ForecastTotal-WorkBudget,0)&" — review Budget")')
    ws[f"{d}{r2}"] = '=IF(ContUsed>0.75,AsOf,"")'
    ws[f"{p}{r2}"] = f'=IF({d}{r2}="","",1)'
    ws[f"{t}{r2}"] = (f'=IF({p}{r2}="","","Contingency "&ROUND(ContUsed*100,0)&"% used — '
                      f'pause optional changes")')
    for rr in (r1, r2):
        ws[f"{d}{rr}"].number_format = DATE

    # top strip: working budget vs sum of category budgets
    put(ws, "B4", '="Working budget ("&CurSym&")"', "label", bold=True)
    put(ws, "C4", "=WorkBudget", "autogrey", MONEY, bold=True)
    put(ws, "E4", "Sum of category budgets", "label", bold=True)
    ws.merge_cells("E4:F4")
    put(ws, "G4", "=BudgetSum", "autogrey", MONEY, bold=True)
    put(ws, "H4", '=IF(WorkBudget=0,"",IF(ABS(BudgetSum-WorkBudget)<1,"✓ matches",'
        'IF(BudgetSum>WorkBudget,"▲ "&FIXED(BudgetSum-WorkBudget,0)&" more than working budget",'
        '"▼ "&FIXED(WorkBudget-BudgetSum,0)&" not yet allocated")))', "label", color=C.GREY_TEXT)
    ws.merge_cells("H4:L4")
    C.level_cf(ws, f"{T.L('status')}{T.first}:{T.L('status')}{tot}", f"${T.L('lvl')}{T.first}")
    ws.freeze_panes = f"C{T.first}"
    C.protect(ws)


# ======================================================================= QUOTES
def build_quotes(ctx):
    ws = ctx.ws(C.S_QUOTES)
    ws.sheet_properties.tabColor = C.TAB_GREEN
    C.title(ws, "⚖️ Quotes",
            "What this tab does: compares up to 4 contractors per trade side by side and "
            "reveals the TRUE cost once hidden extras are added.")
    C.widths(ws, {"B": 30, "C": 17, "D": 12, "E": 17, "F": 12, "G": 17, "H": 12, "I": 17,
                  "J": 12, "K": 3, "L": 3})
    section(ws, "B4", "SCORING WEIGHTS (points, should total 100)", "H")
    wnames = [("C", "Price", "WPrice", 40), ("D", "Timeline", "WTime", 20),
              ("E", "Warranty", "WWar", 15), ("F", "Checklist", "WComp", 15),
              ("G", "Gut & refs", "WGut", 10)]
    for L, lab, name, v in wnames:
        put(ws, f"{L}5", lab, "header")
        put(ws, f"{L}6", v, "input", INT, align="center")
        C.define(ctx.wb, name, C.S_QUOTES, f"${L}$6")
    put(ws, "H5", "Total", "header")
    put(ws, "H6", "=SUM(C6:G6)", "autogrey", INT, align="center", bold=True)
    C.add_cf(ws, "H6", "H6<>100", C.BAD, bold=True)
    put(ws, "B6", "Cheapest TRUE cost, shortest timeline and longest warranty get full points; "
        "others get proportional points.", "note", wrap=True)
    ws.row_dimensions[6].height = 30
    put(ws, "B7", "Tip: type prices exactly as quoted. Yellow checklist cells: estimate what "
        "a missing item will cost you — it is added to the TRUE COST.", "note")

    blocks = {b["cat"]: b for b in ctx.data.QUOTES} if ctx.demo else {}
    demo_blocks = list(ctx.data.QUOTES) if ctx.demo else []
    labels = {
        "name": "Contractor / company", "price": '="Quoted price ("&CurSym&")"',
        "incl": '="Price includes "&TaxWord&"?"', "norm": '="Price incl. "&TaxWord&" ("&CurSym&")"',
        "weeks": "Timeline (weeks)", "warranty": "Warranty (years)",
        "refs": "References checked?", "gut": "Gut feeling (1–5)", "valid": "Quote valid until",
        "extras": '="Hidden extras ("&CurSym&")"', "true": '="TRUE COST ("&CurSym&")"',
        "complete": "Completeness", "p_price": "Price points", "p_time": "Timeline points",
        "p_war": "Warranty points", "p_comp": "Completeness points",
        "p_gut": "Gut & references points", "total": "TOTAL SCORE", "rank": "Ranking",
    }
    notes = {
        "incl": "Yes = the price already includes tax. No = Renovation OS adds your default "
                "tax rate (Start Here) so quotes compare like-for-like.",
        "gut": "1 = bad feeling, 5 = great. Trust matters on a 6-month job.",
        "valid": "Expiry date written on the quote. You get an alert before it expires "
                 "(unless the category is already Contracted).",
        "true": "Price incl. tax + all estimated hidden extras. This is what the job will "
                "really cost you.",
        "complete": "Share of the 13 checklist items marked Included.",
    }
    for k in range(QUOTE_BLOCKS):
        b = qblock(k)
        R = {key: b + off for key, off in QO.items()}
        blk = demo_blocks[k] if k < len(demo_blocks) else None
        # header band
        for L in "BCDEFGHIJ":
            style(ws[f"{L}{b}"], "header", align="left")
        ws[f"B{b}"].value = f"COMPARISON {k + 1}  ▸  trade:"
        cat = ws[f"C{b}"]
        cat.value = blk["cat"] if blk else None
        style(cat, "input", bold=True)
        ctx.dv.add(ws, "Categories", f"C{b}")
        ws.merge_cells(f"E{b}:J{b}")
        ws[f"E{b}"].value = (
            f'=IF(COUNT($C${R["total"]}:$J${R["total"]})=0,"Enter up to 4 quotes below",'
            f'"🏆 Best overall: "&INDEX($C${R["name"]}:$J${R["name"]},MATCH(MAX($C${R["total"]}:'
            f'$J${R["total"]}),$C${R["total"]}:$J${R["total"]},0))&"  —  true cost "&CurSym&" "&'
            f'FIXED(INDEX($C${R["true"]}:$J${R["true"]},MATCH(MAX($C${R["total"]}:$J${R["total"]}),'
            f'$C${R["total"]}:$J${R["total"]},0)),0))')
        ws.row_dimensions[b].height = 28
        # sub header
        for j, (vc, xc) in enumerate(zip(QUOTE_VCOLS, QUOTE_XCOLS)):
            put(ws, f"{vc}{R['sub']}", f"Contractor {j + 1}", "band", align="center")
            ws.merge_cells(f"{vc}{R['sub']}:{xc}{R['sub']}")
        # label column
        for key, lab in labels.items():
            c = put(ws, f"B{R[key]}", lab, "label",
                    bold=key in ("true", "total", "name", "rank"))
            if key in notes:
                C.comment(c, notes[key])
        put(ws, f"B{R['chk_head']}", "HIDDEN-COST CHECKLIST", "band")
        put(ws, f"B{R['score_head']}", "SCORE", "band")
        for i, item in enumerate(CHECKLIST):
            lab = '=TaxWord&" shown in price"' if i == 8 else item
            put(ws, f"B{R['chk_first'] + i}", lab, "label")
        c0, c1 = R["chk_first"], R["chk_first"] + len(CHECKLIST) - 1
        for j, (vc, xc) in enumerate(zip(QUOTE_VCOLS, QUOTE_XCOLS)):
            cd = blk["contractors"][j] if blk and j < len(blk["contractors"]) else None
            put(ws, f"{vc}{R['chk_head']}", "Included?", "band", align="center")
            put(ws, f"{xc}{R['chk_head']}", '="Extra "&CurSym', "band", align="center")
            # inputs
            inputs = [("name", None, None), ("price", MONEY, None), ("incl", None, "YesNo"),
                      ("weeks", "0.#", None), ("warranty", "0.#", None), ("refs", None, "YesNo"),
                      ("gut", INT, "Gut"), ("valid", DATE, None)]
            for key, fmt, lst in inputs:
                v = cd.get(key) if cd else None
                if key == "incl" and v is None:
                    v = "=DefInclTax"
                put(ws, f"{vc}{R[key]}", v, "input", fmt, align="center")
                if lst:
                    ctx.dv.add(ws, lst, f"{vc}{R[key]}")
            for i in range(len(CHECKLIST)):
                st, ex = ("Included", None) if cd else (None, None)
                if cd and i in cd.get("chk", {}):
                    st, ex = cd["chk"][i]
                put(ws, f"{vc}{c0 + i}", st, "input", align="center")
                put(ws, f"{xc}{c0 + i}", ex, "input", MONEY)
            ctx.dv.add(ws, "Incl", f"{vc}{c0}:{vc}{c1}")
            v = vc
            n = R
            f = {
                "norm": f'=IF({v}{n["price"]}="","",IF({v}{n["incl"]}="No",{v}{n["price"]}*(1+VatRate),{v}{n["price"]}))',
                "extras": f'=IF({v}{n["name"]}="","",SUMPRODUCT(({v}{c0}:{v}{c1}<>"Included")*{xc}{c0}:{xc}{c1}))',
                "true": f'=IF(OR({v}{n["name"]}="",{v}{n["norm"]}=""),"",{v}{n["norm"]}+{v}{n["extras"]})',
                "complete": f'=IF({v}{n["name"]}="","",COUNTIF({v}{c0}:{v}{c1},"Included")/{len(CHECKLIST)})',
                "p_price": f'=IF({v}{n["true"]}="","",IF({v}{n["true"]}>0,MIN($C{n["true"]}:$J{n["true"]})/{v}{n["true"]},0)*WPrice)',
                "p_time": f'=IF({v}{n["true"]}="","",IF({v}{n["weeks"]}>0,MIN($C{n["weeks"]}:$J{n["weeks"]})/{v}{n["weeks"]},0)*WTime)',
                "p_war": f'=IF({v}{n["true"]}="","",IF(MAX($C{n["warranty"]}:$J{n["warranty"]})>0,{v}{n["warranty"]}/MAX($C{n["warranty"]}:$J{n["warranty"]}),0)*WWar)',
                "p_comp": f'=IF({v}{n["true"]}="","",{v}{n["complete"]}*WComp)',
                "p_gut": f'=IF({v}{n["true"]}="","",(IF({v}{n["gut"]}="",0,{v}{n["gut"]}/5)*0.7+IF({v}{n["refs"]}="Yes",0.3,0))*WGut)',
                "total": f'=IF({v}{n["true"]}="","",ROUND(SUM({v}{n["p_price"]}:{v}{n["p_gut"]}),1))',
                "rank": f'=IF({v}{n["total"]}="","",IF({v}{n["total"]}=MAX($C{n["total"]}:$J{n["total"]}),"🏆 WINNER","#"&(1+COUNTIF($C{n["total"]}:$J{n["total"]},">"&{v}{n["total"]}))))',
            }
            fmts = {"norm": MONEY, "extras": MONEY, "true": MONEY, "complete": PCT,
                    "p_price": "0.0", "p_time": "0.0", "p_war": "0.0", "p_comp": "0.0",
                    "p_gut": "0.0", "total": "0.0", "rank": None}
            for key, fml in f.items():
                kind = "band" if key in ("true", "total") else "autogrey"
                put(ws, f"{v}{n[key]}", fml, kind, fmts[key], align="center",
                    bold=key in ("true", "total", "rank"))
            # merge value + extra column on single-value rows
            for key in ("name", "price", "incl", "norm", "weeks", "warranty", "refs", "gut",
                        "valid", "extras", "true", "complete", "p_price", "p_time", "p_war",
                        "p_comp", "p_gut", "total", "rank"):
                ws.merge_cells(f"{vc}{R[key]}:{xc}{R[key]}")
        # checklist conditional formats (one rule per block, relative refs)
        sq = " ".join(f"{vc}{c0}:{vc}{c1}" for vc in QUOTE_VCOLS)
        C.add_cf(ws, sq, f'C{c0}="Not included"', C.BAD)
        C.add_cf(ws, sq, f'C{c0}="Unclear"', C.WARN)
        C.add_cf(ws, sq, f'C{c0}="Included"', C.GOOD)
        sqx = " ".join(f"{xc}{c0}:{xc}{c1}" for xc in QUOTE_XCOLS)
        C.add_cf_fill(ws, sqx, f'C{c0}="Included"', C.AUTO_BG, C.GREY_TEXT)
        sqr = " ".join(f"{vc}{R['rank']}" for vc in QUOTE_VCOLS)
        C.add_cf(ws, sqr, f'AND(C{R["total"]}<>"",C{R["total"]}=MAX($C{R["total"]}:$J{R["total"]}))',
                 C.GOOD, bold=True)
        # warning row
        w = R["warn"]
        ws.merge_cells(f"B{w}:J{w}")
        nm, nr, tr = R["name"], R["norm"], R["true"]
        put(ws, f"B{w}",
            f'=IF(COUNT($C{tr}:$J{tr})<2,"",IF(INDEX($C{nm}:$J{nm},MATCH(MIN($C{nr}:$J{nr}),$C{nr}:$J{nr},0))'
            f'<>INDEX($C{nm}:$J{nm},MATCH(MIN($C{tr}:$J{tr}),$C{tr}:$J{tr},0)),'
            f'"⚠️ Hidden costs change the ranking! Cheapest quoted price: "&INDEX($C{nm}:$J{nm},'
            f'MATCH(MIN($C{nr}:$J{nr}),$C{nr}:$J{nr},0))&"  →  cheapest TRUE cost: "&'
            f'INDEX($C{nm}:$J{nm},MATCH(MIN($C{tr}:$J{tr}),$C{tr}:$J{tr},0)),'
            f'"✓ The cheapest quote is also the cheapest true cost."))',
            "auto", bold=True, align="center")
        ws[f"P{w}"] = f'=IF(B{w}="",0,IF(LEFT(B{w},1)="✓",0,1))'
        C.add_cf(ws, f"B{w}", f"$P{w}=1", C.BAD, bold=True)
        C.add_cf(ws, f"B{w}", f'AND($P{w}=0,B{w}<>"")', C.GOOD)
        ws.row_dimensions[w].height = 26
        # hidden per-block summary for Budget "best quote"
        s = QSUM_FIRST + k
        ws[f"M{s}"] = f'=IF(C{b}="","",C{b})'
        ws[f"N{s}"] = f'=IF(OR(M{s}="",COUNT($C{tr}:$J{tr})=0),"",MIN($C{tr}:$J{tr}))'
        ws[f"O{s}"] = f'=IF(N{s}="",0,IF(N{s}>0,1/N{s},0))'
        # hidden quote-expiry alerts
        stage = (f'IFERROR(INDEX({BUDGET.a("stage")},MATCH($C${b},{BUDGET.a("cat")},0)),"")')
        for j, vc in enumerate(QUOTE_VCOLS):
            e = QEXP_FIRST + k * 4 + j
            nmc, vd = f"{vc}{R['name']}", f"{vc}{R['valid']}"
            ws[f"T{e}"] = (f'=IF(OR({nmc}="",{vd}=""),"",IF(OR({stage}="Contracted",'
                           f'{stage}="Complete"),"",IF({vd}<AsOf,"",{vd})))')
            ws[f"T{e}"].number_format = DATE
            ws[f"U{e}"] = f"={pri_from_due(f'T{e}')}"
            ws[f"V{e}"] = (f'=IF(U{e}="","","Quote expires: "&{nmc}&IF($C${b}="",""," ("&$C${b}&")")'
                           f'&" — decide or ask for an extension")')
    for L, h in (("M", "block cat"), ("N", "best true"), ("O", "1/best"), ("P", "warn"),
                 ("T", "expiry due"), ("U", "pri"), ("V", "text")):
        ws[f"{L}8"] = h
        ws.column_dimensions[L].hidden = True
    for L in ("Q", "R", "S"):
        ws.column_dimensions[L].hidden = True
    C.define(ctx.wb, "QBlockCat", C.S_QUOTES, f"$M${QSUM_FIRST}:$M${QSUM_FIRST + QUOTE_BLOCKS - 1}")
    C.define(ctx.wb, "QBlockInv", C.S_QUOTES, f"$O${QSUM_FIRST}:$O${QSUM_FIRST + QUOTE_BLOCKS - 1}")
    ws.freeze_panes = "C8"
    C.protect(ws)


# ======================================================================= CONTRACTORS
def build_contractors(ctx):
    ws = ctx.ws(C.S_CON)
    ws.sheet_properties.tabColor = C.TAB_GREEN
    C.title(ws, "👷 Contractors",
            "What this tab does: one row per contractor — contract, paperwork, progress and "
            "a warning when you pay ahead of the work. Communication log below.")
    T = CON
    spec = [
        dict(key="company", header="Company", width=26, kind="in"),
        dict(key="trade", header="Trade", width=20, kind="in"),
        dict(key="cat", header="Budget category", width=17, kind="in", list="Categories",
             note="Links this contract to a Budget line."),
        dict(key="room", header="Main room", width=15, kind="in", list="Rooms",
             note="Where most of their work happens (or Whole house). Used by the Rooms tab."),
        dict(key="contact", header="Contact person", width=16, kind="in"),
        dict(key="phone", header="Phone", width=15, kind="in"),
        dict(key="email", header="Email", width=24, kind="in"),
        dict(key="value", header=cur_hdr("Contract value"), width=14, kind="in", fmt=MONEY,
             note="Signed contract amount incl. tax (without later change orders)."),
        dict(key="co_add", header=cur_hdr("Approved changes"), width=14, kind="auto",
             fmt=MONEY, grey=True,
             f=f'=IF({{company}}="","",SUMIFS({CO.a("cost")},{CO.a("contractor")},{{company}},'
               f'{CO.a("status")},"Approved"))'),
        dict(key="total", header=cur_hdr("Total contract"), width=14, kind="auto", fmt=MONEY,
             bold=True, f='=IF({company}="","",{value}+{co_add})'),
        dict(key="signed", header="Contract signed?", width=11, kind="in", list="YesNo",
             align="center", note="Only signed contracts count as Committed on the Budget."),
        dict(key="start", header="Start", width=13, kind="in", fmt=DATE),
        dict(key="end", header="Expected completion", width=13, kind="in", fmt=DATE),
        dict(key="insured", header="Insurance ✓", width=11, kind="in", list="YesNo",
             align="center", note="Have you seen a valid liability insurance certificate?"),
        dict(key="licensed", header="License ✓", width=11, kind="in", list="YesNo",
             align="center", note="Registered / licensed for this trade where required?"),
        dict(key="warranty", header="Warranty (years)", width=10, kind="in", fmt="0.#",
             align="center"),
        dict(key="work", header="% work complete", width=11, kind="in", fmt=PCT, align="center",
             note="YOUR estimate of how much of their work is done. Update weekly — it drives "
                  "the overpayment check."),
        dict(key="paid", header=cur_hdr("Paid to date"), width=13, kind="auto", fmt=MONEY,
             grey=True,
             f=f'=IF({{company}}="","",SUMIFS({PAY.a("paid_amt")},{PAY.a("party")},{{company}}))'),
        dict(key="paid_pct", header="Paid %", width=9, kind="auto", fmt=PCT, align="center",
             f='=IF(OR({company}="",{total}=0),"",{paid}/{total})'),
        dict(key="flag", header="Overpayment check", width=24, kind="auto", align="center",
             note="⚠️ appears when paid % is ahead of work complete % by more than your "
                  "tolerance (Start Here).",
             f=f'=IF({{lvl}}="","",IF({{lvl}}=2,"{C.FLAG_OVERPAY}","✓ OK"))'),
        dict(key="lvl", header="lvl", kind="hid",
             f='=IF({paid_pct}="","",IF({paid_pct}-{work}>OverpayTol,2,0))'),
        dict(key="due", header="due", kind="hid", fmt=DATE, f='=IF({lvl}=2,AsOf,"")'),
        dict(key="pri", header="pri", kind="hid", f='=IF({due}="","",1)'),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","","Paying ahead of work: "&{company}&" — paid "&ROUND({paid_pct}*100,0)'
               '&"%, work done "&ROUND({work}*100,0)&"%")'),
    ]
    data = ctx.data.CONTRACTORS if ctx.demo else []
    write_table(ws, T, spec, data, ctx.dv)
    C.level_cf(ws, f"{T.L('flag')}{T.first}:{T.L('flag')}{T.last}", f"${T.L('lvl')}{T.first}")
    for key in ("insured", "licensed"):
        L = T.L(key)
        C.add_cf(ws, f"{L}{T.first}:{L}{T.last}", f'{L}{T.first}="No"', C.BAD, bold=True)
        C.add_cf(ws, f"{L}{T.first}:{L}{T.last}", f'{L}{T.first}="Yes"', C.GOOD)

    # communication log
    M = COMM
    section(ws, f"B{M.header_row - 2}", "COMMUNICATION LOG — every call, email and site "
            "conversation. Follow-ups turn into alerts.", "K")
    spec2 = [
        dict(key="who", header="Contractor", kind="in", list="Contractors", warn=True),
        dict(key="date", header="Date", kind="in", fmt=DATE),
        dict(key="summary", header="What was said / agreed", kind="in", wrap=True),
        dict(key="fu", header="Follow-up needed?", kind="in", list="YesNo", align="center"),
        dict(key="fu_date", header="Follow-up date", kind="in", fmt=DATE),
        dict(key="done", header="Follow-up done?", kind="in", list="YesNo", align="center"),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f='=IF(AND({who}<>"",{fu}="Yes",{done}<>"Yes",{fu_date}<>""),{fu_date},"")'),
        dict(key="pri", header="pri", kind="hid", f=f"={pri_from_due('{due}')}"),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","","Follow up with "&{who}&": "&{summary})'),
    ]
    write_table(ws, M, spec2, ctx.data.COMM_LOG if ctx.demo else [], ctx.dv)
    for r in M.rows():
        ws.merge_cells(f"D{r}:H{r}")
    ws.merge_cells(f"D{M.header_row}:H{M.header_row}")
    for L in ("V", "W", "X", "Y"):
        ws.column_dimensions[L].hidden = True
    ws.freeze_panes = f"C{T.first}"
    C.protect(ws)


# ======================================================================= PAYMENTS
def build_payments(ctx):
    ws = ctx.ws(C.S_PAY)
    ws.sheet_properties.tabColor = C.TAB_GREEN
    C.title(ws, "💳 Payments",
            "What this tab does: your payment schedule, invoices and payments per "
            "contractor — plus a 12-month cash-flow forecast (right side →).")
    T = PAY
    con_val = f'IFERROR(INDEX({CON.a("value")},MATCH({{party}},{CON.a("company")},0)),0)'
    spec = [
        dict(key="party", header="Paid to (contractor / supplier)", width=24, kind="in",
             list="Contractors", warn=True,
             note="Pick a contractor, or type any supplier name (you'll get a warning you can "
                  "accept)."),
        dict(key="milestone", header="Milestone / what for", width=26, kind="in"),
        dict(key="room", header="Room (optional)", width=14, kind="in", list="Rooms",
             note="Leave empty to use the contractor's main room."),
        dict(key="cat_in", header="Category (if not a contractor)", width=16, kind="in",
             list="Categories",
             note="Only needed for suppliers that are not on the Contractors tab, so the "
                  "payment lands on the right Budget line."),
        dict(key="pct", header="% of contract", width=9, kind="in", fmt=PCT, align="center",
             note="Milestone as % of the contract value…"),
        dict(key="fixed", header=cur_hdr("…or fixed amount"), width=12, kind="in", fmt=MONEY,
             note="…or a fixed amount (wins over %)."),
        dict(key="sched", header=cur_hdr("Scheduled"), width=12, kind="auto", fmt=MONEY,
             grey=True,
             f='=IF({party}="","",IF({fixed}<>"",{fixed},IF({pct}<>"",{pct}*'
               + con_val + ',"")))'),
        dict(key="due_date", header="Due date", width=13, kind="in", fmt=DATE),
        dict(key="inv_no", header="Invoice #", width=13, kind="in"),
        dict(key="inv_date", header="Invoice date", width=13, kind="in", fmt=DATE),
        dict(key="inv_amt", header=cur_hdr("Invoiced"), width=12, kind="in", fmt=MONEY),
        dict(key="paid_date", header="Paid date", width=13, kind="in", fmt=DATE,
             note="Enter the date you paid — this marks the row as Paid."),
        dict(key="paid_amt", header=cur_hdr("Paid"), width=12, kind="in", fmt=MONEY),
        dict(key="status", header="Status", width=12, kind="auto", align="center", bold=True,
             f='=IF({party}="","",IF({paid_date}<>"","Paid",IF(AND({due_date}<>"",{due_date}<AsOf),'
               '"OVERDUE",IF(OR({inv_date}<>"",{inv_amt}<>""),"Invoiced","Scheduled"))))'),
        dict(key="cat", header="Category (auto)", width=16, kind="auto", grey=True,
             f='=IF({party}="","",IF({cat_in}<>"",{cat_in},IFERROR(INDEX('
               + CON.a("cat") + ',MATCH({party},' + CON.a("company") + ',0)),"")))'),
        dict(key="room_auto", header="room", kind="hid",
             f='=IF({party}="","",IF({room}<>"",{room},IFERROR(INDEX('
               + CON.a("room") + ',MATCH({party},' + CON.a("company") + ',0)),"")))'),
        dict(key="lvl", header="lvl", kind="hid",
             f='=IF({status}="","",IF({status}="OVERDUE",2,IF({status}="Paid",0,'
               'IF(AND({due_date}<>"",{due_date}<=AsOf+AlertDays),1,0))))'),
        dict(key="unpaid", header="unpaid", kind="hid",
             f='=IF(AND({party}<>"",{paid_date}="",{due_date}<>""),IF({inv_amt}<>"",{inv_amt},'
               'IF({sched}="",0,{sched})),0)'),
        dict(key="inv_due", header="1/due", kind="hid",
             f='=IF(AND({party}<>"",{paid_date}="",{due_date}<>""),1/{due_date},0)'),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f='=IF(AND({party}<>"",{paid_date}="",{due_date}<>""),{due_date},"")'),
        dict(key="pri", header="pri", kind="hid", f=f"={pri_from_due('{due}')}"),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","",IF({status}="OVERDUE","Overdue payment: ","Payment due: ")&{party}'
               '&" — "&{milestone}&" ("&CurSym&" "&FIXED({unpaid},0)&")")'),
    ]
    write_table(ws, T, spec, ctx.data.PAYMENTS if ctx.demo else [], ctx.dv)
    S = T.L("status")
    rng = f"{S}{T.first}:{S}{T.last}"
    C.add_cf(ws, rng, f'{S}{T.first}="OVERDUE"', C.BAD, bold=True)
    C.add_cf(ws, rng, f'{S}{T.first}="Paid"', C.GOOD)
    C.add_cf(ws, rng, f'AND({S}{T.first}<>"",${T.L("lvl")}{T.first}=1)', C.WARN)
    ws.column_dimensions["Y"].width = 3

    # per-contractor summary
    P = PAYSUM
    section(ws, "Z4", "PER-CONTRACTOR SUMMARY", "AF")
    pl = lambda k: T.local(k)  # noqa: E731
    spec2 = [
        dict(key="name", header="Contractor", width=24, kind="auto", bold=True,
             f=f'=IF({C.q(C.S_CON)}!{CON.L("company")}{{r}}="","",{C.q(C.S_CON)}!{CON.L("company")}{{r}})'),
        dict(key="contract", header=cur_hdr("Total contract"), width=13, kind="auto", fmt=MONEY,
             f=f'=IF({{name}}="","",{C.q(C.S_CON)}!{CON.L("total")}{{r}})'),
        dict(key="invoiced", header=cur_hdr("Invoiced"), width=12, kind="auto", fmt=MONEY,
             f=f'=IF({{name}}="","",SUMIFS({pl("inv_amt")},{pl("party")},{{name}}))'),
        dict(key="paid", header=cur_hdr("Paid"), width=12, kind="auto", fmt=MONEY,
             f=f'=IF({{name}}="","",SUMIFS({pl("paid_amt")},{pl("party")},{{name}}))'),
        dict(key="outstanding", header=cur_hdr("Still to pay"), width=12, kind="auto",
             fmt=MONEY, bold=True, f='=IF({name}="","",{contract}-{paid})'),
        dict(key="next_date", header="Next payment", width=13, kind="auto", fmt=DATE,
             f=f'=IF({{name}}="","",IFERROR(1/SUMPRODUCT(MAX(({pl("party")}={{name}})*{pl("inv_due")})),""))'),
        dict(key="next_amt", header=cur_hdr("Next amount"), width=12, kind="auto", fmt=MONEY,
             f=f'=IF({{next_date}}="","",SUMIFS({pl("unpaid")},{pl("party")},{{name}},'
               f'{pl("due_date")},{{next_date}}))'),
    ]
    write_table(ws, P, spec2)

    # cash flow
    section(ws, "Z37", "CASH FLOW — 12 months from project start", "AF")
    put(ws, "Z38", '="Cash needed in the next 30 days ("&CurSym&")"', "label", bold=True)
    ws.merge_cells("Z38:AB38")
    c = put(ws, CASH30_CELL, f'=SUMIFS({pl("unpaid")},{pl("due_date")},"<="&(AsOf+30))', "band",
            MONEY, bold=True, size=12)
    C.comment(c, "All unpaid amounts due in the next 30 days, including anything overdue.")
    C.define(ctx.wb, "Cash30", C.S_PAY, "$" + CASH30_CELL[:2] + "$" + CASH30_CELL[2:])
    K = CASH
    hdrs = {"month": "Month", "c_paid": cur_hdr("Paid"), "c_sched": cur_hdr("Still to pay"),
            "c_total": cur_hdr("Total"), "c_cum": cur_hdr("Cumulative")}
    for key, h in hdrs.items():
        put(ws, f"{K.L(key)}{K.header_row}", h, "header")
    for i, r in enumerate(K.rows()):
        m = K.c("month", r)
        nxt = f"DATE(YEAR({m}),MONTH({m})+1,1)"
        put(ws, m, f'=IF(ProjStart="","",DATE(YEAR(ProjStart),MONTH(ProjStart)+{i},1))', "auto",
            "mmm yyyy", bold=True)
        put(ws, K.c("c_paid", r), f'=IF({m}="","",SUMIFS({pl("paid_amt")},{pl("paid_date")},'
            f'">="&{m},{pl("paid_date")},"<"&{nxt}))', "auto", MONEY)
        put(ws, K.c("c_sched", r), f'=IF({m}="","",SUMIFS({pl("unpaid")},{pl("due_date")},'
            f'">="&{m},{pl("due_date")},"<"&{nxt}))', "auto", MONEY)
        put(ws, K.c("c_total", r), f'=IF({m}="","",{K.c("c_paid", r)}+{K.c("c_sched", r)})',
            "auto", MONEY, bold=True)
        put(ws, K.c("c_cum", r), f'=IF({m}="","",SUM(${K.L("c_total")}${K.first}:{K.c("c_total", r)}))',
            "autogrey", MONEY)
        ws.row_dimensions[r].height = 20
    C.define(ctx.wb, "CashMonths", C.S_PAY, f"${K.L('month')}${K.first}:${K.L('month')}${K.last}")
    ws.freeze_panes = f"C{T.first}"
    C.protect(ws)


# ======================================================================= CHANGE ORDERS
def build_change_orders(ctx):
    ws = ctx.ws(C.S_CO)
    ws.sheet_properties.tabColor = C.TAB_GREEN
    C.title(ws, "🔁 Change Orders",
            "What this tab does: logs every change to the plan and previews its effect on "
            "budget, contingency and move-in date BEFORE you approve it.")
    ws.merge_cells("B4:Q4")
    c = put(ws, "B4", "⚠️  Never approve a change without checking this preview  →  the grey "
            "columns show your forecast, contingency and move-in date IF you approve.",
            "auto", bold=True, size=11, align="center")
    c.fill = C.fill(C.BAD[0])
    c.font = C.font(11, True, C.BAD[1])
    ws.row_dimensions[4].height = 30
    T = CO
    spec = [
        dict(key="no", header="CO #", width=8, kind="in", align="center"),
        dict(key="date", header="Date requested", width=13, kind="in", fmt=DATE),
        dict(key="req", header="Requested by", width=14, kind="in", list="ReqBy"),
        dict(key="contractor", header="Contractor", width=22, kind="in", list="Contractors",
             warn=True),
        dict(key="cat", header="Budget category", width=16, kind="in", list="Categories"),
        dict(key="room", header="Room", width=14, kind="in", list="Rooms"),
        dict(key="desc", header="What changes", width=34, kind="in", wrap=True),
        dict(key="reason", header="Reason", width=18, kind="in", list="COReason"),
        dict(key="cost", header=cur_hdr("Extra cost"), width=12, kind="in", fmt=MONEY,
             note="Extra cost incl. tax. Use a negative number for savings."),
        dict(key="days", header="Extra days", width=9, kind="in", fmt=INT, align="center",
             note="How many days this adds. After approving, also add these days as Delay "
                  "on the affected Timeline task."),
        dict(key="status", header="Status", width=11, kind="in", list="COStatus",
             align="center", default=None),
        dict(key="dec_date", header="Decision date", width=13, kind="in", fmt=DATE),
        dict(key="new_total", header=cur_hdr("Forecast if approved"), width=14, kind="auto",
             fmt=MONEY, grey=True, bold=True,
             note="Budget forecast total including this change (already included if Approved).",
             f='=IF(AND({desc}="",{cost}=""),"",ForecastTotal+IF({status}="Approved",0,'
               'IF({cost}="",0,{cost})))'),
        dict(key="cont_left", header=cur_hdr("Contingency left if approved"), width=14,
             kind="auto", fmt=MONEY, grey=True,
             f='=IF({new_total}="","",ContAmt-MAX(0,{new_total}-WorkBudget))'),
        dict(key="new_end", header="Move-in if approved", width=14, kind="auto", fmt=DATE,
             grey=True,
             f='=IF(OR({new_total}="",ForecastEnd=""),"",ForecastEnd+IF({status}="Approved",0,'
               'IF({days}="",0,{days})))'),
        dict(key="pct_over", header="vs working budget", width=11, kind="auto",
             fmt='+0%;-0%;0%', grey=True, align="center",
             f='=IF(OR({new_total}="",WorkBudget=0),"",{new_total}/WorkBudget-1)'),
        dict(key="lvl", header="lvl", kind="hid",
             f='=IF({new_total}="","",IF({new_total}<=WorkBudget,0,IF({new_total}<=TotalBudget,1,2)))'),
        dict(key="due", header="due", kind="hid", fmt=DATE,
             f='=IF(AND(OR({desc}<>"",{cost}<>""),OR({status}="Pending",{status}=""),{date}<>""),'
               '{date}+3,"")'),
        dict(key="pri", header="pri", kind="hid", f='=IF({due}="","",IF(AsOf>{due},1,2))'),
        dict(key="text", header="text", kind="hid",
             f='=IF({pri}="","","Decide on change order "&{no}&": "&{desc}&" (+"&CurSym&" "&'
               'FIXED(IF({cost}="",0,{cost}),0)&")")'),
    ]
    write_table(ws, T, spec, ctx.data.CHANGE_ORDERS if ctx.demo else [], ctx.dv, row_height=24)
    lv = f"${T.L('lvl')}{T.first}"
    C.level_cf(ws, f"{T.L('new_total')}{T.first}:{T.L('pct_over')}{T.last}", lv)
    S = T.L("status")
    rng = f"{S}{T.first}:{S}{T.last}"
    C.add_cf(ws, rng, f'{S}{T.first}="Pending"', C.WARN, bold=True)
    C.add_cf(ws, rng, f'{S}{T.first}="Approved"', C.GOOD)
    C.add_cf_fill(ws, rng, f'{S}{T.first}="Rejected"', C.AUTO_BG, C.GREY_TEXT)
    ws.freeze_panes = f"D{T.first}"
    C.protect(ws)
