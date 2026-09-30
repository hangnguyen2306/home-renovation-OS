"""Overview tabs: Engine (hidden alert brain), Dashboard, This Week."""
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment

from . import core as C
from .core import DATE, INT, MONEY, PCT, put, section, style
from .layout import (BUDGET, CASH, CO, ENGINE_FIRST, ISS, PAY, SEL, TL, TOP_N, CON,
                     engine_blocks)


# ======================================================================= ENGINE
def build_engine(ctx):
    ws = ctx.ws(C.S_ENG)
    ws.sheet_properties.tabColor = C.TAB_GREY
    hdr = ["Source tab", "Alert", "Due", "Priority", "Sort key", "Week key"]
    for i, h in enumerate(hdr):
        c = ws.cell(row=1, column=i + 1, value=h)
        c.font = C.font(10, True)
    blocks = engine_blocks(ctx.lite)
    last = ENGINE_FIRST + sum(b[3] for b in blocks) - 1
    E_RNG = lambda col: f"${col}${ENGINE_FIRST}:${col}${last}"  # noqa: E731
    r = ENGINE_FIRST
    for label, sheet, first, n, tcol, dcol, pcol in blocks:
        qs = C.q(sheet)
        for sr in range(first, first + n):
            ws[f"A{r}"] = label
            ws[f"B{r}"] = f"={qs}!{tcol}{sr}"
            ws[f"C{r}"] = f"={qs}!{dcol}{sr}"
            ws[f"C{r}"].number_format = DATE
            ws[f"D{r}"] = f"={qs}!{pcol}{sr}"
            ws[f"E{r}"] = (f'=IF(OR(B{r}="",C{r}="",D{r}=""),"",D{r}*100000+(C{r}-DATE(2000,1,1))'
                           f'+ROW()/100000)')
            if not ctx.lite:   # week ranking only feeds This Week (Pro)
                ws[f"F{r}"] = f'=IF(E{r}="","",IF(C{r}>WeekEndDate,"",E{r}))'
            r += 1
    assert r - 1 == last
    # ranked outputs
    outputs = [("H", "E", "TOP 25 (all)")]
    if not ctx.lite:
        outputs.append(("P", "F", "TOP 25 (due by week end)"))
    for base, key_col, title in outputs:
        cols = [C.col_add(base, i) for i in range(7)]
        for L, h in zip(cols, ["Rank", "Key", "Row", "Tab", "Alert", "Due", "Priority"]):
            ws[f"{L}1"] = h
            ws[f"{L}1"].font = C.font(10, True)
        for k in range(1, TOP_N + 1):
            rr = k + 1
            rk, key, row, tab, txt, due, pri = [f"{L}{rr}" for L in cols]
            ws[rk] = k
            ws[key] = f'=IFERROR(SMALL({E_RNG(key_col)},{rk}),"")'
            ws[row] = f'=IF({key}="","",MATCH({key},{E_RNG(key_col)},0))'
            ws[tab] = f'=IF({row}="","",INDEX({E_RNG("A")},{row}))'
            ws[txt] = f'=IF({row}="","",INDEX({E_RNG("B")},{row}))'
            ws[due] = f'=IF({row}="","",INDEX({E_RNG("C")},{row}))'
            ws[due].number_format = DATE
            ws[pri] = f'=IF({row}="","",INDEX({E_RNG("D")},{row}))'
        ws[f"{base}28"] = title
    stats = [("AlertCount", f"=COUNT({E_RNG('E')})"),
             ("OverdueCount", f'=COUNTIFS({E_RNG("C")},"<"&AsOf,{E_RNG("D")},">=1")'),
             ("UrgentCount", f"=COUNTIF({E_RNG('D')},1)")]
    if not ctx.lite:
        stats.append(("WeekAlertCount", f"=COUNT({E_RNG('F')})"))
    for i, (name, f) in enumerate(stats):
        ws[f"X{i + 2}"] = name
        ws[f"Y{i + 2}"] = f
        C.define(ctx.wb, name, C.S_ENG, f"$Y${i + 2}")
    # chart data: top 12 budget categories by max(budget, forecast)
    ws["AA1"], ws["AB1"], ws["AC1"], ws["AD1"] = "Rank", "Category", "Budget", "Forecast"
    for k in range(1, 13):
        rr = k + 1
        ws[f"AA{rr}"] = k
        ws[f"AB{rr}"] = (f'=IFERROR(INDEX({BUDGET.a("cat")},MATCH(LARGE({BUDGET.a("chartkey")},AA{rr}),'
                         f'{BUDGET.a("chartkey")},0)),"")')
        ws[f"AC{rr}"] = (f'=IF(AB{rr}="","",SUMIFS({BUDGET.a("budget")},{BUDGET.a("cat")},AB{rr}))')
        ws[f"AD{rr}"] = (f'=IF(AB{rr}="","",SUMIFS({BUDGET.a("forecast")},{BUDGET.a("cat")},AB{rr}))')
    ws.sheet_state = "hidden"
    C.protect(ws)


def _title(text):
    """Chart title with a modest font size (openpyxl default is huge)."""
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
    cp = CharacterProperties(sz=1100, b=True, solidFill=C.ACCENT)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


E = lambda col, k: f"{C.q(C.S_ENG)}!{col}{k + 1}"  # noqa: E731


def _when(due):
    return (f'=IF({due}="","",IF({due}<AsOf,(AsOf-{due})&IF(AsOf-{due}=1," day"," days")&" overdue",'
            f'IF({due}=AsOf,"today","in "&({due}-AsOf)&IF({due}-AsOf=1," day"," days"))))')


def _pri_label(p):
    return f'=IF({p}="","",IF({p}=1,"🔴 Urgent",IF({p}=2,"🟡 This week","🔵 Upcoming")))'


def _pri_cf(ws, rng, pcell):
    C.add_cf(ws, rng, f"{pcell}=1", C.BAD, bold=True)
    C.add_cf(ws, rng, f"{pcell}=2", C.WARN, bold=True)
    C.add_cf(ws, rng, f"{pcell}=3", ("DBEAFE", "1E3A8A"))


# ======================================================================= DASHBOARD
def build_dashboard(ctx):
    ws = ctx.ws(C.S_DASH)
    ws.sheet_properties.tabColor = C.TAB_TEAL
    ws.sheet_view.showGridLines = False
    C.title(ws, "", "What this tab does: your renovation at a glance — updated automatically. "
            "Nothing to type here.")
    ws["B1"].value = '="🏠 "&IF(ProjName="","My renovation",ProjName)'
    cols = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]
    C.widths(ws, {L: 13.5 for L in cols})
    C.widths(ws, {"L": 3, "M": 3})
    put(ws, "I2", "As of", "label", color=C.GREY_TEXT, align="right")
    put(ws, "J2", "=AsOf", "label", DATE, bold=True, color=C.ACCENT)

    bar = lambda x: f'REPT("█",ROUND(MIN(1,MAX(0,{x}))*10,0))&REPT("░",10-ROUND(MIN(1,MAX(0,{x}))*10,0))'  # noqa
    cards = [
        [('="Working budget ("&CurSym&")"', "=WorkBudget", MONEY,
          '="of "&FIXED(TotalBudget,0)&" total budget"'),
         ('="Committed ("&CurSym&")"', "=CommittedTotal", MONEY,
          '=IF(WorkBudget=0,"",ROUND(CommittedTotal/WorkBudget*100,0)&"% of working budget")'),
         ('="Invoiced ("&CurSym&")"', "=InvoicedTotal", MONEY,
          '=IF(CommittedTotal=0,"",ROUND(InvoicedTotal/CommittedTotal*100,0)&"% of committed")'),
         ('="Paid ("&CurSym&")"', "=PaidTotal", MONEY,
          '="Cash needed next 30 days: "&FIXED(Cash30,0)'),
         ('="Remaining to commit ("&CurSym&")"', "=WorkBudget-CommittedTotal", MONEY,
          '="working budget − committed"')],
        [('="Forecast final cost ("&CurSym&")"', "=ForecastTotal", MONEY,
          '="incl. risk: "&FIXED(ForecastRisk,0)'),
         ('="Over / under budget ("&CurSym&")"', "=ForecastTotal-WorkBudget", "+#,##0;-#,##0;0",
          '=IF(ForecastTotal>WorkBudget,"▲ over working budget",IF(ForecastTotal<WorkBudget,'
          '"▼ under working budget","exactly on budget"))'),
         ("Contingency used", "=ContUsed", PCT, f"={bar('ContUsed')}"),
         ("Project day", '=IF(OR(ProjStart="",TargetEnd=""),"–","Day "&MAX(0,AsOf-ProjStart+1)&'
          '" of "&(TargetEnd-ProjStart+1))', None,
          '=IF(TargetEnd="","set dates on Start Here",IF(TargetEnd>=AsOf,(TargetEnd-AsOf)&'
          '" days to target move-in","target date has passed"))'),
         ("Time elapsed", '=IF(OR(ProjStart="",TargetEnd=""),"–",IF(TargetEnd<=ProjStart,"–",'
          'MAX(0,MIN(1,(AsOf-ProjStart)/(TargetEnd-ProjStart)))))', PCT, None)],
        [("Forecast move-in", '=IF(ForecastEnd="","–",ForecastEnd)', DATE,
          '=IF(OR(ForecastEnd="",TargetEnd=""),"",IF(ForecastEnd>TargetEnd,(ForecastEnd-TargetEnd)'
          '&" days after target",IF(ForecastEnd<TargetEnd,(TargetEnd-ForecastEnd)&'
          '" days before target","on target")))'),
         ("Open issues", "=SUM(OpenIssues)", INT, '=OpenHigh&" high severity"'),
         ("Overdue items", "=OverdueCount", INT, '="of "&AlertCount&" alerts in total"'),
         ("Decisions waiting", "=DecisionsWaiting", INT, '="see Selections & Orders"'),
         ("Orders due", "=OrdersDue+OrdersLate", INT, '=OrdersLate&" already late"')],
    ]
    if ctx.lite:   # no Selections & Orders tab: show timeline + payment pressure instead
        cards[2][3] = ("Late tasks", f"=COUNTIF({TL.a('code')},3)", INT, '="see Timeline"')
        cards[2][4] = ("Payments due soon", f"=COUNTIF({PAY.a('lvl')},1)", INT,
                       '="within "&AlertDays&" days"')
    tops = [4, 8, 12]
    value_cells = {}
    for ri, row in enumerate(cards):
        t = tops[ri]
        for ci, (lab, val, fmt, sub) in enumerate(row):
            c0, c1 = cols[ci * 2], cols[ci * 2 + 1]
            for rr in (t, t + 1, t + 2):
                ws.merge_cells(f"{c0}{rr}:{c1}{rr}")
                for L in (c0, c1):
                    cell = ws[f"{L}{rr}"]
                    cell.fill = C.fill(C.ACCENT_LIGHT)
                    cell.protection = C.LOCKED
            lc = ws[f"{c0}{t}"]
            lc.value = lab
            lc.font = C.font(9, True, C.GREY_TEXT)
            lc.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
            vc = ws[f"{c0}{t + 1}"]
            vc.value = val
            vc.font = C.font(18, True, C.ACCENT)
            vc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            if fmt:
                vc.number_format = fmt
            sc = ws[f"{c0}{t + 2}"]
            if ri == 1 and ci == 4:
                x = f"{c0}{t + 1}"
                sub = (f'=IF({x}="–","",REPT("█",ROUND({x}*10,0))&'
                       f'REPT("░",10-ROUND({x}*10,0)))')
            sc.value = sub
            sc.font = C.font(8.5, False, C.GREY_TEXT)
            sc.alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=False)
            value_cells[(ri, ci)] = f"{c0}{t + 1}"
        ws.row_dimensions[t].height = 18
        ws.row_dimensions[t + 1].height = 30
        ws.row_dimensions[t + 2].height = 18
        ws.row_dimensions[t + 3].height = 8
    # KPI colouring
    v = value_cells
    C.add_cf(ws, v[(1, 1)], f"{v[(1, 1)]}>0", C.BAD, bold=True)
    C.add_cf(ws, v[(1, 1)], f"{v[(1, 1)]}<=0", C.GOOD, bold=True)
    C.add_cf(ws, v[(1, 2)], f"{v[(1, 2)]}>0.75", C.BAD, bold=True)
    C.add_cf(ws, v[(1, 2)], f"AND({v[(1, 2)]}>0.5,{v[(1, 2)]}<=0.75)", C.WARN, bold=True)
    C.add_cf(ws, v[(2, 0)], f'AND({v[(2, 0)]}<>"–",TargetEnd<>"",{v[(2, 0)]}>TargetEnd)', C.BAD,
             bold=True)
    C.add_cf(ws, v[(2, 2)], f"{v[(2, 2)]}>0", C.BAD, bold=True)
    C.add_cf(ws, v[(2, 1)], "OpenHigh>0", C.BAD, bold=True)
    if ctx.lite:
        C.add_cf(ws, v[(2, 3)], f"{v[(2, 3)]}>0", C.BAD, bold=True)
    else:
        C.add_cf(ws, v[(2, 3)], f"{v[(2, 3)]}>0", C.WARN, bold=True)
        C.add_cf(ws, v[(2, 4)], "OrdersLate>0", C.BAD, bold=True)
    C.add_cf(ws, v[(2, 4)], f"{v[(2, 4)]}>0", C.WARN, bold=True)

    # ------------------------------------------------------------ HEALTH
    section(ws, "B16", "RENOVATION HEALTH", "K")
    ws.merge_cells("B17:K17")
    put(ws, "B17", '=IF(N17=0,"🟢  ON TRACK",IF(N17=1,"🟡  NEEDS ATTENTION","🔴  ACTION REQUIRED"))',
        "auto", size=16, bold=True, align="center")
    ws.row_dimensions[17].height = 36
    ws["N17"] = "=MAX(N18:N23)"
    C.level_cf(ws, "B17", "$N$17")
    tl = TL
    helpers = {
        19: (f'=COUNTIFS({tl.a("code")},3,{tl.a("affects")},">=1")', None, None),
        20: (f'=COUNTIF({PAY.a("status")},"OVERDUE")', f"=COUNTIF({CON.a('lvl')},2)",
             f"=COUNTIF({PAY.a('lvl')},1)"),
        22: (f"=SUM({SEL.a('dec_over')})", f"=SUM({SEL.a('dec_soon')})", None),
    }
    rows = [
        (18, "Budget", '=IF(ForecastTotal<=WorkBudget,0,IF(ForecastTotal<=TotalBudget,1,2))',
         '="Forecast "&CurSym&" "&FIXED(ForecastTotal,0)&" vs working budget "&CurSym&" "&'
         'FIXED(WorkBudget,0)&IF(N18=0,"  — within budget",IF(N18=1,"  — eating into contingency",'
         '"  — above your TOTAL budget!"))'),
        (19, "Timeline", '=IF(O19>0,2,IF(OR(ForecastEnd="",TargetEnd=""),0,IF(ForecastEnd>TargetEnd+14,'
         '2,IF(ForecastEnd>TargetEnd,1,0))))',
         '=IF(ForecastEnd="","Add a project start date and tasks",IF(O19>0,O19&" late task(s) '
         'that other tasks depend on · ","")&"forecast move-in "&IF(TargetEnd="","",IF(ForecastEnd>'
         'TargetEnd,(ForecastEnd-TargetEnd)&" days after target","on or before target")))'),
        (20, "Contractors & payments", '=IF(OR(O20>0,P20>0),2,IF(Q20>0,1,0))',
         '=O20&" overdue payment(s) · "&P20&" paying-ahead warning(s) · "&Q20&" payment(s) due '
         'within "&AlertDays&" days"'),
        (21, "Materials", '=IF(OrdersLate+DeliveriesLate>0,2,IF(OrdersDue>0,1,0))',
         '=OrdersLate&" late order(s) · "&DeliveriesLate&" delivery after install date · "&'
         'OrdersDue&" order(s) due within "&AlertDays&" days"'),
        (22, "Decisions", '=IF(O22>0,2,IF(P22>0,1,0))',
         '=O22&" decision(s) past decide-by date · "&P22&" due within "&AlertDays&" days"'),
        (23, "Issues", '=IF(OpenHigh>0,2,IF(SUM(OpenIssues)>0,1,0))',
         '=OpenHigh&" open high-severity · "&SUM(OpenIssues)&" open in total"'),
    ]
    if ctx.lite:   # no Materials / Decisions rows: Issues moves up, rows 22-23 are hidden
        del helpers[22]
        issues = rows[5]
        rows = rows[:3] + [(21,) + issues[1:]]
        for r in (22, 23):
            ws.row_dimensions[r].hidden = True
    for r, name, lvl, reason in rows:
        ws.merge_cells(f"B{r}:C{r}")
        put(ws, f"B{r}", name, "label", bold=True)
        put(ws, f"D{r}", f'=IF(N{r}=0,"🟢",IF(N{r}=1,"🟡","🔴"))', "auto", align="center", size=12)
        ws.merge_cells(f"E{r}:K{r}")
        put(ws, f"E{r}", reason, "label", color=C.DARK)
        ws[f"N{r}"] = lvl
        for j, L in enumerate(("O", "P", "Q")):
            if r in helpers and helpers[r][j]:
                ws[f"{L}{r}"] = helpers[r][j]
        C.level_cf(ws, f"D{r}", f"$N${r}")
        ws.row_dimensions[r].height = 24
        for L in "BCDEFGHIJK":
            ws[f"{L}{r}"].border = C.Border(bottom=C.Side(style="thin", color=C.LINE))

    # ------------------------------------------------------------ NEXT ACTIONS
    section(ws, "B25", "YOUR NEXT ACTIONS — ranked by urgency", "K")
    ws.merge_cells("B26:K26")
    put(ws, "B26", '=IF(AlertCount=0,"🟢  Nothing needs your attention right now",IF(UrgentCount>0,'
        '"🔴  ","🟡  ")&AlertCount&IF(AlertCount=1," ACTION NEEDS"," ACTIONS NEED")&" YOUR ATTENTION")',
        "auto", size=14, bold=True, align="center")
    ws["N26"] = '=IF(AlertCount=0,0,IF(UrgentCount>0,2,1))'
    C.level_cf(ws, "B26", "$N$26")
    ws.row_dimensions[26].height = 32
    hdr = [("B", "B", "#"), ("C", "C", "Priority"), ("D", "G", "What to do"),
           ("H", "I", "Where"), ("J", "J", "Due"), ("K", "K", "When")]
    for a, b, h in hdr:
        if a != b:
            ws.merge_cells(f"{a}27:{b}27")
        put(ws, f"{a}27", h, "header")
    for k in range(1, 11):
        r = 27 + k
        ws[f"O{r}"] = f"={E('N', k)}"
        put(ws, f"B{r}", f'=IF({E("L", k)}="","",{k})', "auto", INT, align="center",
            color=C.GREY_TEXT)
        put(ws, f"C{r}", _pri_label(f"O{r}"), "auto", align="center", size=9)
        ws.merge_cells(f"D{r}:G{r}")
        put(ws, f"D{r}", f'=IF({E("L", k)}="","",{E("L", k)})', "auto", wrap=True, size=9)
        ws.merge_cells(f"H{r}:I{r}")
        put(ws, f"H{r}", f'=IF({E("K", k)}="","",{E("K", k)})', "auto", color=C.GREY_TEXT)
        put(ws, f"J{r}", f'=IF({E("M", k)}="","",{E("M", k)})', "auto", DATE, align="center")
        put(ws, f"K{r}", _when(f"J{r}"), "auto", align="center", size=9)
        ws.row_dimensions[r].height = 30
    _pri_cf(ws, "C28:C37", "$O28")
    C.add_cf(ws, "K28:K37", '$O28=1', C.BAD, bold=True)
    more = "each tab" if ctx.lite else "This Week and in each tab"
    put(ws, "B38", f'=IF(AlertCount>10,"+ "&(AlertCount-10)&" more — see {more}.","")', "note")

    # ------------------------------------------------------------ CHARTS
    section(ws, "B40", "CHARTS", "K")
    eng = ctx.ws(C.S_ENG)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.title = _title("Budget vs forecast (top 12 categories)")
    ch.y_axis.number_format = "#,##0"
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    data = Reference(eng, min_col=29, max_col=30, min_row=1, max_row=13)
    cats = Reference(eng, min_col=28, min_row=2, max_row=13)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties = GraphicalProperties(solidFill="CBD5E1")
    ch.series[1].graphicalProperties = GraphicalProperties(solidFill=C.ACCENT)
    ch.height, ch.width = 10, 12.6
    ch.gapWidth = 40
    ch.legend.position = "b"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, "B41")

    pay = ctx.ws(C.S_PAY)
    cf = BarChart()
    cf.type = "col"
    cf.grouping = "stacked"
    cf.overlap = 100
    cf.title = _title("Cash flow per month")
    k0 = C.col_idx(CASH.L("month"))
    data = Reference(pay, min_col=k0 + 1, max_col=k0 + 2, min_row=CASH.header_row,
                     max_row=CASH.last)
    cats = Reference(pay, min_col=k0, min_row=CASH.first, max_row=CASH.last)
    cf.add_data(data, titles_from_data=True)
    cf.set_categories(cats)
    cf.series[0].graphicalProperties = GraphicalProperties(solidFill=C.ACCENT)
    cf.series[1].graphicalProperties = GraphicalProperties(solidFill="F59E0B")
    cf.y_axis.number_format = "#,##0"
    cf.y_axis.majorGridlines = None
    cf.x_axis.number_format = "mmm yy"
    cf.height, cf.width = 10, 12.6
    cf.gapWidth = 60
    cf.legend.position = "b"
    cf.x_axis.delete = False
    cf.y_axis.delete = False
    ws.add_chart(cf, "G41")
    for L in ("N", "O", "P", "Q"):
        ws.column_dimensions[L].hidden = True
    C.protect(ws)


# ======================================================================= THIS WEEK
def build_this_week(ctx):
    ws = ctx.ws(C.S_WEEK)
    ws.sheet_properties.tabColor = C.TAB_TEAL
    ws.sheet_view.showGridLines = False
    C.title(ws, "🗓️ This Week",
            "What this tab does: Monday — see everything that falls in this week. Friday — "
            "review what happened and write down next week's top 3.")
    C.widths(ws, {"B": 30, "C": 14, "D": 46, "E": 22, "F": 14, "G": 16, "H": 3})
    put(ws, "B4", "Week starting", "label", bold=True, size=11)
    put(ws, "C4", '=AsOf-IF(WeekStart="Sunday",WEEKDAY(AsOf,1)-1,WEEKDAY(AsOf,3))', "input", DATE,
        bold=True, align="center")
    C.comment(ws["C4"], "Shows the current week automatically. Type any date to plan another "
              "week (to go back, type  =AsOf  or re-download the formula from the guide).")
    put(ws, "B5", "Week ending", "label", color=C.GREY_TEXT)
    put(ws, "C5", '=C4+6', "autogrey", DATE, align="center")
    put(ws, "D4", "Tip: overwrite the yellow date to look at another week.", "note")
    C.define(ctx.wb, "WeekStartDate", C.S_WEEK, "$C$4")
    C.define(ctx.wb, "WeekEndDate", C.S_WEEK, "$C$5")
    wk = lambda rng: f'{rng},">="&WeekStartDate,{rng},"<="&WeekEndDate'  # noqa: E731

    section(ws, "B7", "📋 MONDAY PLAN — what falls in this week", "G")
    plan = [
        ("Decisions to make", f"=COUNTIFS({wk(SEL.a('decide_by'))},{SEL.a('dec_wait')},1)",
         "Selections with a decide-by date this week"),
        ("Orders to place", f"=COUNTIFS({wk(SEL.a('order_by'))},{SEL.a('need_order')},1)",
         "Chosen items with an order-by date this week"),
        ("Payments due", f'=COUNTIFS({wk(PAY.a("due_date"))},{PAY.a("unpaid")},">0")',
         "Unpaid payments with a due date this week"),
        ("Deliveries expected", f'=COUNTIFS({wk(SEL.a("delivery"))},{SEL.a("delivered")},"<>Yes")',
         "Orders arriving this week — someone must be on site"),
        ("Tasks starting", f"=COUNTIFS({wk(TL.a('start'))})", "Timeline tasks starting"),
        ("Tasks finishing", f"=COUNTIFS({wk(TL.a('end'))})", "Timeline tasks due to finish"),
        ("Issues with an action this week", f"=COUNTIFS({wk(ISS.a('act_date'))},{ISS.a('open')},1)",
         '="Open issues in total: "&SUM(OpenIssues)'),
    ]
    for i, (lab, f, note) in enumerate(plan):
        r = 8 + i
        put(ws, f"B{r}", lab, "label", bold=True)
        put(ws, f"C{r}", f, "band", INT, size=12, bold=True, align="center")
        put(ws, f"D{r}", note, "note")
        C.add_cf(ws, f"C{r}", f"C{r}>0", C.WARN, bold=True)
        ws.row_dimensions[r].height = 22

    section(ws, "B16", "🔴 TOP 10 ALERTS DUE BY THE END OF THIS WEEK (incl. overdue)", "G")
    for L, h in zip("BCDEF", ["Priority", "Due", "What to do", "Where", "When"]):
        put(ws, f"{L}17", h, "header")
    ws.merge_cells("F17:G17")
    for k in range(1, 11):
        r = 17 + k
        ws[f"I{r}"] = f"={E('V', k)}"
        put(ws, f"B{r}", _pri_label(f"I{r}"), "auto", align="center", size=9)
        put(ws, f"C{r}", f'=IF({E("U", k)}="","",{E("U", k)})', "auto", DATE, align="center")
        put(ws, f"D{r}", f'=IF({E("T", k)}="","",{E("T", k)})', "auto", wrap=True, size=9)
        put(ws, f"E{r}", f'=IF({E("S", k)}="","",{E("S", k)})', "auto", color=C.GREY_TEXT)
        ws.merge_cells(f"F{r}:G{r}")
        put(ws, f"F{r}", _when(f"C{r}"), "auto", align="center", size=9)
        ws.row_dimensions[r].height = 30
    _pri_cf(ws, "B18:B27", "$I18")
    ws.column_dimensions["I"].hidden = True
    put(ws, "B28", '=IF(WeekAlertCount=0,"🟢 Nothing due this week.",IF(WeekAlertCount>10,"+ "&'
        '(WeekAlertCount-10)&" more due this week","All "&WeekAlertCount&" shown."))', "note")

    section(ws, "B30", "📝 FRIDAY REVIEW — the numbers (automatic)", "G")
    rev = [
        ('="Paid this week ("&CurSym&")"', f"=SUMIFS({PAY.a('paid_amt')},{wk(PAY.a('paid_date'))})",
         MONEY),
        ("Change orders approved this week",
         f'=COUNTIFS({CO.a("status")},"Approved",{wk(CO.a("dec_date"))})', INT),
        ('="…their extra cost ("&CurSym&")"',
         f'=SUMIFS({CO.a("cost")},{CO.a("status")},"Approved",{wk(CO.a("dec_date"))})', MONEY),
        ("Tasks completed this week", f'=COUNTIFS({TL.a("status")},"Done",{wk(TL.a("done_date"))})',
         INT),
        ("Issues opened this week", f"=COUNTIFS({wk(ISS.a('found'))})", INT),
        ("Issues closed this week", f"=COUNTIFS({wk(ISS.a('res_date'))})", INT),
    ]
    for i, (lab, f, fmt) in enumerate(rev):
        r = 31 + i
        put(ws, f"B{r}", lab, "label", bold=True)
        put(ws, f"C{r}", f, "band", fmt, size=12, bold=True, align="center")
        ws.row_dimensions[r].height = 22

    section(ws, "B38", "📝 FRIDAY REVIEW — your notes (type in the yellow boxes)", "G")
    notes = ctx.data.FRIDAY_NOTES if ctx.demo else {}
    boxes = [("completed", "What was completed?"), ("delayed", "What was delayed?"),
             ("changed", "What changed?"), ("risk", "What is now at risk?"),
             ("next", "Top 3 for next week")]
    r = 39
    for key, q in boxes:
        put(ws, f"B{r}", q, "label", bold=True, color=C.ACCENT, size=11)
        ws.merge_cells(f"B{r + 1}:G{r + 3}")
        c = ws[f"B{r + 1}"]
        c.value = notes.get(key)
        style(c, "input", wrap=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        for rr in range(r + 1, r + 4):
            for L in "BCDEFG":
                ws[f"{L}{rr}"].fill = C.fill(C.INPUT)
                ws[f"{L}{rr}"].protection = C.UNLOCKED
        r += 5
    C.protect(ws)
