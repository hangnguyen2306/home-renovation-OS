"""Lists (hidden) and Start Here."""
from openpyxl.styles import Alignment

from . import core as C
from .core import put, section, style
from .layout import BUDGET, CON, ROOMDATA

CURRENCIES = [("EUR €", "€"), ("USD $", "$"), ("GBP £", "£"), ("CHF", "CHF"), ("CAD $", "C$"),
              ("AUD $", "A$"), ("NZD $", "NZ$"), ("SEK kr", "kr"), ("NOK kr", "kr"),
              ("DKK kr", "kr"), ("PLN zł", "zł"), ("CZK Kč", "Kč"), ("Other", "")]
TAX_NAMES = ["VAT", "Sales tax", "GST", "HST", "TVA", "MwSt", "BTW", "IVA", "None"]
DOC_TYPES = ["Contract", "Quote", "Invoice", "Permit", "Plan", "Insurance", "Warranty",
             "Receipt", "Photo", "Inspection", "Other"]


def register_lists(ctx):
    L = ctx.lists
    L.register("YesNo", ["Yes", "No"])
    L.register("Stage", ["Estimating", "Contracted", "Complete"])
    L.register("Incl", ["Included", "Not included", "Unclear"])
    L.register("COStatus", ["Pending", "Approved", "Rejected"])
    L.register("COReason", ["Client choice", "Unforeseen condition", "Design change",
                            "Code requirement", "Contractor error"])
    L.register("ReqBy", ["Owner", "Contractor", "Architect / designer", "Inspector", "Other"])
    L.register("Phase", ["Pre-construction", "Construction", "Final"])
    L.register("TaskStatus", ["Not started", "In progress", "Done"])
    L.register("IssueType", ["Issue", "Punch item"])
    L.register("Severity", ["Low", "Medium", "High"])
    L.register("IssueStatus", ["Open", "In progress", "Resolved"])
    L.register("DocType", DOC_TYPES)
    L.register("PhotoStage", ["Before", "Demolition", "Rough-in", "Walls", "Finishing",
                              "Final"])
    L.register("ABC", ["A", "B", "C"])
    L.register("Gut", [1, 2, 3, 4, 5])
    L.register("CurLabel", [c[0] for c in CURRENCIES])
    L.register("CurSym", [c[1] for c in CURRENCIES])
    L.register("TaxName", TAX_NAMES)
    L.register("AreaUnit", ["m²", "ft²"])
    L.register("WeekStart", ["Monday", "Sunday"])
    L.register("WarnBand", [0, 0.05, 0.1])
    L.register("Categories", [
        f'=IF({C.q(C.S_BUDGET)}!{BUDGET.c("cat", r)}="","",{C.q(C.S_BUDGET)}!{BUDGET.c("cat", r)})'
        for r in BUDGET.rows()])
    if ctx.lite:   # no Rooms tab: fixed room list, and any other name may be typed
        L.register("Rooms", [n for n in ctx.data.BLANK_ROOM_NAMES[:6]] + ["Whole house"])
        ctx.dv.force_warn.add("Rooms")
    else:
        L.register("Rooms", [
            f'=IF({C.q(C.S_ROOMS)}!{ROOMDATA.c("name", r)}="","",'
            f'{C.q(C.S_ROOMS)}!{ROOMDATA.c("name", r)})' for r in ROOMDATA.rows()]
            + ["Whole house"])
    L.register("Contractors", [
        f'=IF({C.q(C.S_CON)}!{CON.c("company", r)}="","",{C.q(C.S_CON)}!{CON.c("company", r)})'
        for r in CON.rows()])


def build_lists(ctx):
    ws = ctx.ws(C.S_LISTS)
    for name, (L, values, header) in ctx.lists.items.items():
        ws[f"{L}1"].value = header
        ws[f"{L}1"].font = C.font(10, True)
        ws.column_dimensions[L].width = 18
        for i, v in enumerate(values):
            c = ws[f"{L}{i + 2}"]
            c.value = v
            c.font = C.font()
            if name == "WarnBand":
                c.number_format = C.PCT
    ws.sheet_state = "hidden"
    ws.sheet_properties.tabColor = C.TAB_GREY
    C.protect(ws)


START_INPUTS = [
    # row, name, label, kind, fmt, hint, comment
    (5, "ProjName", "Project name", None, None, "Shown as the Dashboard title.", None),
    (6, "Address", "Property address", None, None, "For your own reference.", None),
    (7, "ProjStart", "Project start date", None, C.DATE,
     "First day of the project (design counts!). The Timeline starts here.", None),
    (8, "TargetEnd", "Target move-in date", None, C.DATE,
     "The date you want to live in the house again.", None),
    (9, "TotalBudget", None, None, C.MONEY,
     "Everything you can spend, including tax and your safety buffer.",
     "Your total renovation budget including tax. Renovation OS keeps the contingency % "
     "aside as a safety buffer, the rest is your working budget."),
    (10, "TotalArea", None, None, C.NUM, "Optional — used for cost per m² / ft².", None),
    (11, "AsOfOverride", "“Today” override (optional)", None, C.DATE,
     "Leave EMPTY to use today's date.",
     "Normally empty. Type a date here to see the workbook as it was/will be on that "
     "date. The demo file is frozen on a date so every feature shows data."),
    (14, "Currency", "Currency", "CurLabel", None,
     "Shown next to every amount header. Pick ‘Other’ to type your own symbol.", None),
    (15, "CustomSym", "Custom currency symbol", None, None, "Only used when Currency = Other.",
     None),
    (16, "TaxName", "Tax name", "TaxName", None, "VAT, Sales tax, GST… used in labels.", None),
    (17, "VatRate", "Default tax rate", None, C.PCT,
     "Used to add tax to quotes entered WITHOUT tax.",
     "When a quote says prices exclude tax, Quotes adds this rate so all quotes are "
     "compared like-for-like."),
    (18, "DefInclTax", "Prices usually quoted incl. tax?", "YesNo", None,
     "Pre-fills each quote; you can still change it per quote.", None),
    (19, "AreaUnit", "Area unit", "AreaUnit", None, "m² or ft² for room sizes.", None),
    (20, "WeekStart", "Week starts on", "WeekStart", None,
     "Used by This Week and the Gantt chart.", None),
    (21, "BudgetWarnPct", "Budget ‘attention’ band", "WarnBand", C.PCT,
     "A budget line turns 🟡 when forecast is over budget by up to this %, 🔴 above.", None),
    (24, "ContPct", "Contingency (safety buffer) %", None, C.PCT,
     "10–15% is typical. Kept aside from your total budget.",
     "Money you keep aside for surprises. Working budget = total budget − contingency."),
    (25, "AlertDays", "Alert window (days)", None, C.INT,
     "Things due within this many days show as 🟡 ‘this week’.", None),
    (26, "DecideBuf", "Decide-before-order buffer (days)", None, C.INT,
     "Decide this many days before the order-by date.",
     "Time you want between making a decision (e.g. choosing tiles) and placing the order."),
    (27, "OrderBuf", "Order safety buffer (days)", None, C.INT,
     "Extra days between delivery and installation.",
     "Order-by = needed-on date − supplier lead time − this buffer."),
    (28, "OverpayTol", "Overpayment tolerance %", None, C.PCT,
     "Warn when paid % is more than this ahead of work done %.",
     "Protects you from paying ahead of progress. Example: 10% means paying 50% for 35% "
     "of the work triggers a warning."),
]


def build_start(ctx):
    ws = ctx.ws(C.S_START)
    wb = ctx.wb
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C.TAB_TEAL
    C.title(ws, "🏠 Renovation OS Lite — Start Here" if ctx.lite else "🏠 Renovation OS — Start Here",
            "What this tab does: set up your project once (yellow cells), then follow the "
            "5 steps on the right.")
    C.widths(ws, {"B": 34, "C": 22, "D": 50, "E": 3, "F": 5, "G": 16, "H": 16, "I": 16,
                  "J": 16, "K": 16, "L": 6})
    section(ws, "B4", "YOUR PROJECT", "D")
    section(ws, "B13", "PREFERENCES", "D")
    section(ws, "B23", "ALERT SETTINGS", "D")
    section(ws, "B30", "CALCULATED FOR YOU", "D")

    d = ctx.data
    s = d.SETUP if ctx.demo else {}
    cur = ctx.cur
    values = {
        "ProjName": s.get("name"), "Address": s.get("address"), "ProjStart": s.get("start"),
        "TargetEnd": s.get("target"), "TotalBudget": s.get("total_budget"),
        "TotalArea": (ctx.area(s.get("area_m2")) if s else None),
        "AsOfOverride": d.AS_OF if ctx.demo else None,
        "Currency": cur["currency"], "CustomSym": None, "TaxName": cur["tax"],
        "VatRate": cur["rate"], "DefInclTax": "Yes", "AreaUnit": cur["area"],
        "WeekStart": cur["week"], "BudgetWarnPct": 0.05, "ContPct": 0.10, "AlertDays": 7,
        "DecideBuf": 7, "OrderBuf": 5, "OverpayTol": 0.10,
    }
    inputs = START_INPUTS
    if ctx.lite:   # these settings only drive Quotes / Selections (Pro): leave them out and
        skip = {"VatRate", "DefInclTax", "DecideBuf", "OrderBuf"}   # close the gaps
        inputs, prev, shift = [], None, 0
        for row, name, *rest in START_INPUTS:
            if prev is not None and row - prev > 2:
                shift = 0   # new section
            prev = row
            if name in skip:
                shift += 1
                continue
            inputs.append((row - shift, name, *rest))
    for row, name, label, lst, fmt, hint, note in inputs:
        lab = ws[f"B{row}"]
        if name == "TotalBudget":
            label = '="Total budget incl. "&TaxWord&" ("&CurSym&")"'
        if name == "TotalArea":
            label = '="Total floor area ("&AreaUnit&")"'
        lab.value = label
        style(lab, "label", bold=True)
        cell = ws[f"C{row}"]
        cell.value = values.get(name)
        style(cell, "input", fmt, align="left")
        if lst:
            ctx.dv.add(ws, lst, f"C{row}")
        h = ws[f"D{row}"]
        h.value = hint
        style(h, "note")
        if note:
            C.comment(lab, note)
        C.define(wb, name, C.S_START, f"$C${row}")
        ws.row_dimensions[row].height = 22
    if ctx.demo:
        ws["D11"].value = ("DEMO is frozen on this date so every feature shows data — "
                           "clear this cell to use today's date.")
        ws["D11"].font = C.font(9, True, C.BAD[1], True)

    calc = [
        (31, "ContAmt", '="Contingency amount ("&CurSym&")"', "=IF(TotalBudget=\"\",0,TotalBudget*ContPct)", C.MONEY),
        (32, "WorkBudget", '="Working budget ("&CurSym&")"', "=IF(TotalBudget=\"\",0,TotalBudget-ContAmt)", C.MONEY),
        (33, "CurSym", "Currency symbol in use",
         f'=IF(Currency="Other",IF(CustomSym="","¤",CustomSym),IFERROR(INDEX({ctx.lists.ref("CurSym")},'
         f'MATCH(Currency,{ctx.lists.ref("CurLabel")},0)),"¤"))', None),
        (34, "AsOf", "Today (as used by all tabs)", '=IF(AsOfOverride="",TODAY(),AsOfOverride)', C.DATE),
        (35, "TaxWord", "Tax word used in labels", '=IF(OR(TaxName="None",TaxName=""),"tax",TaxName)', None),
    ]
    for row, name, label, f, fmt in calc:
        put(ws, f"B{row}", label, "label", bold=True)
        put(ws, f"C{row}", f, "autogrey", fmt, bold=True, align="left")
        C.define(wb, name, C.S_START, f"$C${row}")
        ws.row_dimensions[row].height = 22
    put(ws, "D31", "Kept aside for surprises — not part of your working budget.", "note")
    put(ws, "D32", "What you plan to spend. Budget lines should add up to this.", "note")

    # ---------------------------------------------------------- 5-step guide
    section(ws, "F4", "GET STARTED IN 10 MINUTES", "K")
    steps = [
        ("Fill in the yellow cells on the left", "Project dates, total budget, currency, tax. "
         "Everything else calculates itself."),
        ("Budget tab: set a budget per category", "Rename categories if you like. Aim for the "
         "total to match your working budget."),
        ("Contractors + Quotes: add who you work with", "Compare quotes side by side — the "
         "hidden-cost checklist shows the TRUE price."),
        ("Timeline: set durations and dependencies", "Dates, Gantt chart and order-by dates "
         "for your materials follow automatically."),
        ("Every Monday open Dashboard + This Week", "Do the ranked actions. On Friday fill "
         "the Friday Review. That's the whole routine."),
    ]
    if ctx.lite:
        steps[2] = ("Contractors + Payments: who you pay, and when", "Enter each contract and "
                    "its payment milestones — you get a warning if you pay ahead of the work.")
        steps[3] = ("Timeline: set durations and dependencies", "Dates, the Gantt chart and "
                    "late-task alerts follow automatically.")
        steps[4] = ("Every Monday open the Dashboard", "Do the ranked actions, top to bottom. "
                    "That's the whole routine.")
    r = 5
    for i, (head, body) in enumerate(steps, 1):
        n = ws[f"F{r}"]
        n.value = i
        n.font = C.font(16, True, C.WHITE)
        n.fill = C.fill(C.ACCENT)
        n.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"F{r}:F{r + 1}")
        ws.merge_cells(f"G{r}:K{r}")
        ws.merge_cells(f"G{r + 1}:K{r + 1}")
        put(ws, f"G{r}", head, "band", border=False, size=11)
        put(ws, f"G{r + 1}", body, "label", wrap=True, color=C.GREY_TEXT)
        ws.row_dimensions[r + 1].height = 30
        r += 2

    # ---------------------------------------------------------- colour legend
    section(ws, "F17", "COLOUR LEGEND", "K")
    legend = [
        ("", C.INPUT, "Light-yellow cell = YOU TYPE HERE (or pick from the dropdown)."),
        ("", C.WHITE, "White / grey cell = automatic. Protected so you can't break a formula."),
        (C.ICON_OK, C.GOOD[0], "On track — nothing to do."),
        (C.ICON_WARN, C.WARN[0], "Needs attention soon (within your alert window)."),
        (C.ICON_BAD, C.BAD[0], "Urgent — overdue, over budget or at risk."),
    ]
    for i, (icon, bg, text) in enumerate(legend):
        rr = 18 + i
        sw = ws[f"F{rr}"]
        sw.value = icon
        sw.fill = C.fill(bg)
        sw.border = C.BORDER_ALL
        sw.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"G{rr}:K{rr}")
        put(ws, f"G{rr}", text, "label")
        ws.row_dimensions[rr].height = 22
    put(ws, "G23", "Sheets are protected without a password: Review ▸ Unprotect sheet if you "
        "really need to change a formula.", "note")

    # ---------------------------------------------------------- tab map
    section(ws, "F25", "YOUR TABS", "K")
    tabs = [
        (C.S_DASH, "Your renovation at a glance + ranked next actions."),
        (C.S_WEEK, "Monday plan and Friday review."),
        (C.S_BUDGET, "Budget vs committed vs forecast per category."),
        (C.S_QUOTES, "Compare up to 4 quotes per trade, incl. hidden costs."),
        (C.S_CON, "Contractors, overpayment check and communication log."),
        (C.S_PAY, "Payment schedule, invoices and 12-month cash flow."),
        (C.S_CO, "Change orders with impact preview before you approve."),
        (C.S_TL, "Tasks, dependencies and Gantt chart."),
        (C.S_ROOMS, "One card per room: money, progress, problems."),
        (C.S_SEL, "Decisions and orders — with decide-by and order-by dates."),
        (C.S_ISS, "Problems found on site and the final punch list."),
        (C.S_VAULT, "Links to every document and photo + what is missing."),
        (C.S_WAR, "Warranties and home maintenance after the renovation."),
    ]
    tabs = [t for t in tabs if t[0] in ctx.wb.sheetnames]
    for i, (name, desc) in enumerate(tabs):
        rr = 26 + i
        c = ws[f"F{rr}"]
        c.value = "▸"
        c.font = C.font(10, True, C.ACCENT)
        ws.merge_cells(f"G{rr}:H{rr}")
        link = ws[f"G{rr}"]
        link.value = name
        link.hyperlink = f"#{C.q(name)}!A1"
        link.font = C.font(10, True, C.ACCENT)
        link.font = C.Font(name=C.FONT, size=10, bold=True, color=C.ACCENT, underline="single")
        ws.merge_cells(f"I{rr}:K{rr}")
        put(ws, f"I{rr}", desc, "label", color=C.GREY_TEXT)
    if ctx.lite:
        rr = 26 + len(tabs) + 1
        ws.merge_cells(f"F{rr}:K{rr + 1}")
        up = put(ws, f"F{rr}", "⭐ Renovation OS Pro adds: quote comparison with hidden costs, "
                 "change-order impact preview, decide-by / order-by dates for materials, room "
                 "cards, This Week planner, document vault and warranty tracker.", "band",
                 wrap=True, border=False)
        up.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[rr].height = 22
        ws.row_dimensions[rr + 1].height = 22
    ws.freeze_panes = "A3"
    C.protect(ws)
