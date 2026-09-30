#!/usr/bin/env python3
"""Recalculate every Renovation OS file in headless LibreOffice and verify it.

1. Builds all 4 workbooks (build_renovation_os.py).
2. Static checks: formula whitelist lint, no external links, protection, hidden helper sheets.
3. LibreOffice (Python-UNO) recalculation → every formula cell with an error result
   (#REF!, #VALUE!, #NAME?, #DIV/0!, #N/A, Err:5xx incl. 522/523 circular references) fails.
4. The recalculated copy is saved as .xlsx and re-scanned with openpyxl for error strings.
5. DEMO: key numbers are asserted against reference_model.py (pure Python, from demo_data.py).
   BLANK: no alerts, overall health ON TRACK.

Run:  python3 test_workbook.py            (exit code 0 = all good)
"""
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import traceback
from datetime import date, timedelta

import openpyxl

import build_renovation_os as B
from reference_model import Model
from ros import core as C
from ros.layout import QO, qblock, CO, CON, TL

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(HERE, "dist")
RECALC = os.path.join(DIST, "recalculated")
ERR_STRINGS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
FAILS = []


def check(cond, msg):
    if not cond:
        FAILS.append(msg)
        print("   ✗", msg)
    return cond


def near(a, b, tol=0.51):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


# ---------------------------------------------------------------- LibreOffice / UNO
class Office:
    def __init__(self):
        import uno  # noqa: F401  (system python3-uno)
        self.profile = tempfile.mkdtemp(prefix="lo_profile_")
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            self.port = s.getsockname()[1]
        self.proc = subprocess.Popen(
            ["soffice", "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
             f"-env:UserInstallation=file://{self.profile}",
             f"--accept=socket,host=127.0.0.1,port={self.port};urp;"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        import uno
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(120):
            try:
                self.ctx = resolver.resolve(
                    f"uno:socket,host=127.0.0.1,port={self.port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise RuntimeError("could not connect to LibreOffice")
        self.desktop = self.ctx.ServiceManager.createInstanceWithContext(
            "com.sun.star.frame.Desktop", self.ctx)

    def prop(self, name, value):
        from com.sun.star.beans import PropertyValue
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    def open(self, path):
        import uno
        url = uno.systemPathToFileUrl(os.path.abspath(path))
        doc = self.desktop.loadComponentFromURL(url, "_blank", 0, (self.prop("Hidden", True),))
        doc.calculateAll()
        return doc

    def save_xlsx(self, doc, path):
        import uno
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(path)),
                       (self.prop("FilterName", "Calc MS Excel 2007 XML"),))

    def close(self):
        try:
            self.desktop.terminate()
        except Exception:
            pass
        self.proc.terminate()
        try:
            self.proc.wait(10)
        except Exception:
            self.proc.kill()
        shutil.rmtree(self.profile, ignore_errors=True)


class Doc:
    """Small reader around a recalculated UNO spreadsheet document."""

    def __init__(self, doc):
        self.doc = doc
        self.sheets = doc.Sheets

    def cell(self, sheet, addr):
        return self.sheets.getByName(sheet).getCellRangeByName(addr)

    def name_cell(self, name):
        rng = self.doc.NamedRanges.getByName(name).getReferredCells()
        return rng.getCellByPosition(0, 0)

    def v(self, sheet, addr):
        return self.cell(sheet, addr).getValue()

    def s(self, sheet, addr):
        return self.cell(sheet, addr).getString()

    def nv(self, name):
        return self.name_cell(name).getValue()

    def date(self, sheet, addr):
        n = self.v(sheet, addr)
        return date(1899, 12, 30) + timedelta(days=int(round(n))) if n else None

    def errors(self):
        from com.sun.star.sheet.FormulaResult import ERROR
        out = []
        for i in range(self.sheets.Count):
            sh = self.sheets.getByIndex(i)
            ranges = sh.queryFormulaCells(ERROR)
            for addr in ranges.getRangeAddresses():
                for r in range(addr.StartRow, addr.EndRow + 1):
                    for c in range(addr.StartColumn, addr.EndColumn + 1):
                        cell = sh.getCellByPosition(c, r)
                        code = cell.getError()
                        if code:
                            out.append((sh.Name, f"{C.col_let(c + 1)}{r + 1}", code,
                                        cell.getFormula()[:140]))
        return out


# ---------------------------------------------------------------- static checks
def static_checks(path, lite):
    wb = openpyxl.load_workbook(path)
    lint = C.lint_workbook(wb)
    check(not lint, f"{os.path.basename(path)}: lint problems {lint[:3]}")
    check(wb.sheetnames == (C.LITE_SHEETS if lite else C.SHEET_ORDER), f"sheet order {wb.sheetnames}")
    for ws in wb.worksheets:
        check(ws.protection.sheet, f"{ws.title} not protected")
        check(bool(ws.protection.password), f"{ws.title} has no protection password")
        check(ws.protection.selectLockedCells, f"{ws.title}: locked cells must not be selectable")
        if ws.sheet_state == "visible":
            check(ws["B3"].value == C.LICENSE_LINE, f"{ws.title}: licence line missing")
    check(wb.security is not None and wb.security.lockStructure,
          "workbook structure not locked")
    check(wb.properties.creator == C.AUTHOR, "author metadata")
    shown = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for row in ws.iter_rows()
             for c in row if isinstance(c.value, str) and c.value.startswith("=")
             and c.protection.locked and not c.protection.hidden]
    check(not shown, f"{len(shown)} locked formula cells not hidden, e.g. {shown[:3]}")
    check(wb[C.S_ENG].sheet_state == "hidden" and wb[C.S_LISTS].sheet_state == "hidden",
          "Engine/Lists must be hidden")
    check(not getattr(wb, "_external_links", []), "external links present")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and ".xls" in c.value:
                    check(False, f"external reference {ws.title}!{c.coordinate}")
    # every yellow cell is unlocked, every unlocked cell is yellow
    bad = 0
    for ws in wb.worksheets:
        if ws.sheet_state == "hidden":
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c, openpyxl.cell.cell.MergedCell):
                    continue
                if ws.title == C.S_START and c.coordinate == "F18":
                    continue   # legend swatch showing the input colour
                if c.hyperlink is not None:
                    continue   # navigation links stay clickable
                yellow = c.fill is not None and c.fill.fgColor is not None and \
                    str(c.fill.fgColor.rgb).endswith(C.INPUT)
                if yellow and c.protection.locked:
                    bad += 1
                if not yellow and not c.protection.locked:
                    bad += 1
    check(bad == 0, f"{bad} cells where yellow ≠ unlocked")
    # only input cells may look yellow: no conditional-format highlight in a yellow hue
    import colorsys
    for ws in wb.worksheets:
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                f = rule.dxf.fill if rule.dxf is not None else None
                rgb = str(f.fgColor.rgb)[-6:] if f is not None and f.fgColor is not None else None
                if not rgb or not all(ch in "0123456789ABCDEFabcdef" for ch in rgb):
                    continue
                r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
                h, l_, s_ = colorsys.rgb_to_hls(r, g, b)
                check(not (40 <= h * 360 <= 70 and s_ > 0.3),
                      f"{ws.title} CF {cf.sqref}: yellow highlight #{rgb} looks like an input cell")
    # dropdowns exist
    n_dv = sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets)
    check(n_dv > (20 if lite else 40), f"only {n_dv} data validations")


# ---------------------------------------------------------------- recalc scan
def recalc_scan(office, path):
    doc = office.open(path)
    D = Doc(doc)
    errs = D.errors()
    for e in errs[:15]:
        print("     ", e)
    check(not errs, f"{os.path.basename(path)}: {len(errs)} formula error cell(s)")
    os.makedirs(RECALC, exist_ok=True)
    out = os.path.join(RECALC, os.path.basename(path))
    office.save_xlsx(doc, out)
    wb = openpyxl.load_workbook(out, data_only=True)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERR_STRINGS):
                    n += 1
                    if n <= 5:
                        print("      cached error", ws.title, c.coordinate, c.value)
    check(n == 0, f"{os.path.basename(path)}: {n} error value(s) in recalculated xlsx")
    return doc, D


# ---------------------------------------------------------------- DEMO asserts
def demo_asserts(D, cur, lite=False):
    m = Model(B.PRESETS[cur], lite)
    tag = f"{'LITE ' if lite else ''}DEMO {cur}"
    print(f"   reference: forecast {m.forecast:,.0f}, committed {m.committed:,.0f}, "
          f"alerts {len(m.alerts)}, finish {m.forecast_end}")
    pairs = [("WorkBudget", m.work), ("ContAmt", m.cont), ("CommittedTotal", m.committed),
             ("InvoicedTotal", m.invoiced), ("PaidTotal", m.paid), ("ForecastTotal", m.forecast),
             ("PendingCO", m.pending_co), ("IssuesOpenCost", m.open_cost), ("Cash30", m.cash30),
             ("AlertCount", len(m.alerts)),
             ("OverdueCount", sum(1 for a in m.alerts if a[3] < m.asof)),
             ("UrgentCount", sum(1 for a in m.alerts if a[4] == 1)),
             ("OpenHigh", m.open_high)]
    if not lite:
        pairs += [("DecisionsWaiting", sum(s["dec_wait"] for s in m.sel)),
                  ("OrdersDue", sum(s["order_now"] for s in m.sel)),
                  ("OrdersLate", sum(s["order_late"] for s in m.sel))]
    for name, exp in pairs:
        got = D.nv(name)
        check(near(got, exp), f"{tag}: {name} = {got} expected {exp}")
    check(near(D.nv("ContUsed"), m.cont_used, 0.001),
          f"{tag}: ContUsed {D.nv('ContUsed')} vs {m.cont_used}")
    fe = D.name_cell("ForecastEnd").getValue()
    check(near(fe, (m.forecast_end - date(1899, 12, 30)).days),
          f"{tag}: ForecastEnd {fe} vs {m.forecast_end}")
    # budget lines
    for i, c in enumerate(m.cats):
        r = 8 + i
        for col, key in (("E", "best"), ("I", "committed"), ("J", "invoiced"), ("K", "paid"),
                         ("L", "forecast")):
            got = D.cell(C.S_BUDGET, f"{col}{r}")
            exp = c[key]
            if exp is None:
                check(got.getString() == "", f"{tag}: Budget {col}{r} expected blank")
            else:
                check(near(got.getValue(), exp), f"{tag}: Budget {c['name']} {key} = "
                      f"{got.getValue()} expected {exp}")
    # ranked engine output vs reference
    top = m.alerts[:25]
    for k, (row, tab, text, due, pri) in enumerate(top, 1):
        rr = k + 1
        gt, gx = D.s(C.S_ENG, f"K{rr}"), D.s(C.S_ENG, f"L{rr}")
        gd, gp = D.date(C.S_ENG, f"M{rr}"), D.v(C.S_ENG, f"N{rr}")
        check(gt == tab and gx == text and gd == due and int(gp) == pri,
              f"{tag}: rank {k}: got ({gt} | {gx} | {gd} | {gp}) expected ({tab} | {text} | {due} | {pri})")
    # every alert type fires
    tabs = {a[1] for a in m.alerts}
    want_tabs = ["Payments", "Timeline", "Contractors", "Issues & Punch List", "Budget"]
    want = ["Overdue payment", "Payment due", "Late task", "Starts soon", "Paying ahead",
            "Follow up", "High-severity", "Issue action due", "Punch item", "Forecast is over"]
    if not lite:
        want_tabs += ["Selections & Orders", "Quotes", "Change Orders", "Warranty & Maintenance"]
        want += ["Decision overdue", "Decide:", "Order:", "Order late", "Delivery after",
                 "Quote expires", "Decide on change order", "Warranty expires", "Maintenance",
                 "Contingency"]
    for t in want_tabs:
        check(t in tabs, f"{tag}: no alert from {t}")
    for prefix in want:
        check(any(a[2].startswith(prefix) for a in m.alerts), f"{tag}: no '{prefix}' alert")
    # dashboard top-10 mirrors engine
    check(D.s(C.S_DASH, "D28") == m.alerts[0][2], f"{tag}: Dashboard first action")
    check(D.s(C.S_DASH, "B26").endswith("ACTIONS NEED YOUR ATTENTION") and
          str(len(m.alerts)) in D.s(C.S_DASH, "B26"), f"{tag}: headline {D.s(C.S_DASH, 'B26')}")
    # health
    h = m.health()
    keys = (["overall", "budget", "timeline", "contractors", "issues"] if lite else
            ["overall", "budget", "timeline", "contractors", "materials", "decisions", "issues"])
    for r, key in zip(range(17, 24), keys):
        check(int(D.v(C.S_DASH, f"N{r}")) == h[key],
              f"{tag}: health {key} = {D.v(C.S_DASH, f'N{r}')} expected {h[key]}")
    check("ACTION REQUIRED" in D.s(C.S_DASH, "B17"), f"{tag}: overall banner {D.s(C.S_DASH, 'B17')}")
    # timeline dates
    for i, t in enumerate(m.tasks):
        r = TL.first + i
        check(D.date(C.S_TL, f"J{r}") == t["start"] and D.date(C.S_TL, f"L{r}") == t["end"],
              f"{tag}: task {t['task']} dates {D.date(C.S_TL, f'J{r}')}–{D.date(C.S_TL, f'L{r}')} "
              f"expected {t['start']}–{t['end']}")
    check(D.s(C.S_START, "C33") == m.sym, f"{tag}: currency symbol")
    if lite:
        return
    # quotes: kitchen block
    qb = m.qblocks[0]
    b = qblock(0)
    head = D.s(C.S_QUOTES, f"E{b}")
    check(qb["winner"] in head, f"{tag}: quote winner {head} expected {qb['winner']}")
    warn = D.s(C.S_QUOTES, f"B{b + QO['warn']}")
    check(qb["warn"] and warn.startswith("⚠️ Hidden costs change the ranking"),
          f"{tag}: hidden-cost warning missing: {warn}")
    for j, row in enumerate(qb["rows"]):
        col = "CEGI"[j]
        check(near(D.v(C.S_QUOTES, f"{col}{b + QO['true']}"), row["true"]),
              f"{tag}: true cost {row['name']}")
        got = D.v(C.S_QUOTES, f"{col}{b + QO['total']}")
        check(near(got, row["total"], 0.11), f"{tag}: score {row['name']} {got} vs {row['total']}")
    # overpayment flag
    for i, c in enumerate(m.con.values()):
        got = D.s(C.S_CON, f"{CON.L('flag')}{CON.first + i}")
        if c["lvl"] == 2:
            check(got == C.FLAG_OVERPAY, f"{tag}: overpay flag for {c['company']}: {got}")
        elif c["lvl"] == 0:
            check(got == "✓ OK", f"{tag}: no overpay for {c['company']}: {got}")
    # pending CO preview
    for i, x in enumerate(m.cos):
        if x["status"] != "Pending":
            continue
        r = CO.first + i
        nt = m.forecast + x["cost"]
        check(near(D.v(C.S_CO, f"{CO.L('new_total')}{r}"), nt), f"{tag}: CO new total")
        check(near(D.v(C.S_CO, f"{CO.L('cont_left')}{r}"), m.cont - max(0, nt - m.work)),
              f"{tag}: CO contingency left")
        check(D.date(C.S_CO, f"{CO.L('new_end')}{r}") == m.forecast_end + timedelta(x["days"]),
              f"{tag}: CO new end")
    # this week
    exp_ws = m.asof - timedelta(days=(m.asof.weekday() if B.PRESETS[cur]["week"] == "Monday"
                                      else (m.asof.weekday() + 1) % 7))
    check(D.date(C.S_WEEK, "C4") == exp_ws, f"{tag}: week start {D.date(C.S_WEEK, 'C4')}")


def blank_asserts(D, cur, lite=False):
    tag = f"{'LITE ' if lite else ''}BLANK {cur}"
    check(D.nv("AlertCount") == 0, f"{tag}: AlertCount {D.nv('AlertCount')}")
    check(D.nv("ForecastTotal") == 0, f"{tag}: forecast not 0")
    check(int(D.v(C.S_DASH, "N17")) == 0, f"{tag}: health not green")
    check("ON TRACK" in D.s(C.S_DASH, "B17"), f"{tag}: banner {D.s(C.S_DASH, 'B17')}")
    check("Nothing needs your attention" in D.s(C.S_DASH, "B26"), f"{tag}: headline")
    # buyer starts filling in the blank file: dates, budget, other currency, Sunday weeks
    start = date.today() - timedelta(days=30)
    serial = lambda d_: (d_ - date(1899, 12, 30)).days  # noqa: E731
    D.name_cell("ProjStart").setValue(serial(start))
    D.name_cell("TargetEnd").setValue(serial(start + timedelta(days=200)))
    D.name_cell("TotalBudget").setValue(80000)
    D.name_cell("Currency").setString("Other")
    D.name_cell("CustomSym").setString("R$")
    D.name_cell("WeekStart").setString("Sunday")
    D.doc.calculateAll()
    errs = D.errors()
    for e in errs[:10]:
        print("     ", e)
    check(not errs, f"{tag} (filled in): {len(errs)} formula error cell(s)")
    check(D.nv("ForecastEnd") > serial(start), f"{tag}: prefilled timeline has no dates")
    check(D.nv("AlertCount") >= 1, f"{tag}: expected 'starts soon' alerts after entering dates")
    check(D.s(C.S_BUDGET, "C7") == "Budget (R$)", f"{tag}: custom currency header {D.s(C.S_BUDGET, 'C7')}")
    if not lite:
        check(D.date(C.S_WEEK, "C4").weekday() == 6, f"{tag}: week should start on Sunday")


def main():
    t0 = time.time()
    paths = []
    for ed in ("PRO", "LITE"):
        for cur in ("EUR",):
            for var in ("DEMO", "BLANK"):
                paths.append((ed == "LITE", cur, var, B.build(cur, var, DIST, ed)))
    print(f"built {len(paths)} files in {time.time() - t0:.1f}s")
    office = Office()
    try:
        for lite, cur, var, path in paths:
            print(f"── {os.path.basename(path)}")
            n0 = len(FAILS)
            static_checks(path, lite)
            doc, D = recalc_scan(office, path)
            try:
                if var == "DEMO":
                    demo_asserts(D, cur, lite)
                else:
                    blank_asserts(D, cur, lite)
            except Exception:
                traceback.print_exc()
                FAILS.append(f"{path}: exception during asserts")
            doc.close(True)
            print("   ✓ ok" if len(FAILS) == n0 else f"   {len(FAILS) - n0} failure(s)")
    finally:
        office.close()
    print(f"\n{'ALL CHECKS PASSED' if not FAILS else f'{len(FAILS)} FAILURE(S)'} "
          f"({time.time() - t0:.0f}s)")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
